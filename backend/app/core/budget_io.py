"""Import/export delle voci di costo: template Excel e lettura di .xlsx/.csv con errori riga per riga.

L'import non calcola nulla: legge le celle, le converte nei tipi del modello e restituisce le voci (o gli errori di
validazione con il numero di riga) al frontend, che poi le invia al motore per la validazione.
"""
from __future__ import annotations

import csv
import io
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError

from app.core.demo import build_demo
from app.core.field_catalog import FIELDS, coerce_value
from app.models.schemas import CostItemInput

MAX_BYTES = 2_000_000
MAX_ROWS = 2000
_BY_NAME = {f["name"].lower(): f for f in FIELDS}
_BY_LABEL = {f["label"].lower(): f for f in FIELDS}


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Voci"
    header_fill = PatternFill("solid", fgColor="1F2937")
    label_fill = PatternFill("solid", fgColor="E5E7EB")
    for col, f in enumerate(FIELDS, 1):
        c = ws.cell(row=1, column=col, value=f["name"])
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        lab = ws.cell(row=2, column=col, value=f"{f['label']}" + (f" [crit. {', '.join(map(str, f['criteria']))}]" if f["criteria"] else ""))
        lab.fill = label_fill
        lab.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = 22
    ws.cell(row=2, column=1, value="(etichetta — riga ignorata all'import)")
    ws.row_dimensions[2].height = 48
    examples = [i for i in build_demo(mode="realistic")["cost_items"] if i["item_id"] in ("P-01", "A-01", "C-01", "O-01")]
    for r, item in enumerate(examples, 3):
        for col, f in enumerate(FIELDS, 1):
            v = item.get(f["name"])
            ws.cell(row=r, column=col, value=", ".join(v) if isinstance(v, list) else v)
    ws.freeze_panes = "B3"

    guide = wb.create_sheet("Guida ai campi")
    for col, h in enumerate(["Campo", "Etichetta", "Tipo", "Gruppo", "Categorie", "Criteri", "Valori ammessi / note"], 1):
        c = guide.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
    for r, f in enumerate(FIELDS, 2):
        note = (", ".join(f.get("options", [])) + " · " if f.get("options") else "") + f["help"]
        for col, v in enumerate([f["name"], f["label"], f["type"], f["group"], ", ".join(f["categories"] or ["tutte"]),
                                 ", ".join(map(str, f["criteria"])), note], 1):
            guide.cell(row=r, column=col, value=v)
    for col, w in zip("ABCDEFG", (28, 40, 10, 24, 34, 16, 70)):
        guide.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _rows(filename: str, data: bytes) -> List[List[Any]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb["Voci"] if "Voci" in wb.sheetnames else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
        return [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    raise ValueError("Formato non supportato: usare .xlsx o .csv")


def parse_import(filename: str, data: bytes) -> Dict[str, Any]:
    if len(data) > MAX_BYTES:
        raise ValueError("File troppo grande (max 2 MB)")
    rows = _rows(filename, data)
    if not rows:
        raise ValueError("File vuoto")
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    mapping: Dict[int, Dict[str, Any]] = {}
    unknown: List[str] = []
    for idx, h in enumerate(header):
        f = _BY_NAME.get(h) or _BY_LABEL.get(h)
        if f:
            mapping[idx] = f
        elif h:
            unknown.append(h)
    if not any(f["name"] == "item_id" for f in mapping.values()):
        raise ValueError("Intestazione non riconosciuta: manca la colonna item_id (usare il template scaricabile)")

    items: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    for r_no, row in enumerate(rows[1:MAX_ROWS + 1], start=2):
        first = str(row[0]).strip() if row and row[0] is not None else ""
        if not any(c is not None and str(c).strip() for c in row) or first.startswith(("(", "#")):
            continue
        raw: Dict[str, Any] = {}
        try:
            for idx, f in mapping.items():
                if idx < len(row):
                    v = coerce_value(f, row[idx])
                    if v is not None:
                        raw[f["name"]] = v
            items.append(CostItemInput.model_validate(raw).model_dump(mode="json", exclude_none=True))
        except ValidationError as exc:
            errors.append({"row": r_no, "message": "; ".join(f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors())})
        except ValueError as exc:
            errors.append({"row": r_no, "message": str(exc)})
    return {"items": items, "errors": errors, "rows_read": len(rows) - 1, "ignored_columns": unknown}
