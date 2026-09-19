"""Layer 1 del Dual-Layer Output: esportazione XLSX con CEP-ID e QR verso l'Auditor Portal."""
from __future__ import annotations

import io
import os
from urllib.parse import quote

import qrcode
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill

from app.models.schemas import BudgetValidationResponse


def verification_url(project_id: str, merkle_root: str) -> str:
    base = os.getenv("QUANTO_PUBLIC_URL", "http://localhost:5173").rstrip("/")
    return f"{base}/?tab=auditor&project={quote(project_id)}&root={merkle_root}"


def build_xlsx(resp: BudgetValidationResponse) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget validato"
    bold = Font(bold=True)
    ws["A1"], ws["A2"], ws["A3"], ws["A4"] = "Progetto", "Bando", "CEP-ID", "Merkle Root"
    ws["B1"], ws["B2"], ws["B3"], ws["B4"] = resp.project_id, resp.bando_id, resp.cep_id, resp.merkle_root
    for cell in ("A1", "A2", "A3", "A4"):
        ws[cell].font = bold
    ws["A5"], ws["B5"] = "Conformity Score", f"{resp.conformity_score}/100"
    ws["A5"].font = bold

    header = ["ID", "Descrizione", "Categoria", "Stato", "Richiesto €", "Ammesso €", "Costo orario €/h", "Tetto €/h", "Hash riga (SHA-256)"]
    start = 8
    for col, name in enumerate(header, 1):
        c = ws.cell(row=start, column=col, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
    for r, i in enumerate(resp.items, start + 1):
        for col, v in enumerate([i.item_id, i.description, i.category.value, i.status.value, i.original_cost_eur,
                                 i.computed_cost_eur, i.hourly_rate_computed or None, i.hourly_rate_cap or None, i.item_hash_sha256], 1):
            ws.cell(row=r, column=col, value=v)
        for col in (5, 6, 7, 8):
            ws.cell(row=r, column=col).number_format = "#,##0.00"
    total_row = start + len(resp.items) + 1
    ws.cell(row=total_row, column=4, value="TOTALE").font = bold
    ws.cell(row=total_row, column=5, value=resp.total_requested_eur).number_format = "#,##0.00"
    ws.cell(row=total_row, column=6, value=resp.total_approved_eur).number_format = "#,##0.00"
    for col, width in zip("ABCDEFGHI", (12, 42, 16, 24, 14, 14, 16, 12, 68)):
        ws.column_dimensions[col].width = width
    ws.cell(row=total_row + 2, column=1, value="Verifica ex-post: inquadrare il QR (Auditor Portal) o usare il CEP-ID.").alignment = Alignment(wrap_text=False)

    qr_buf = io.BytesIO()
    qrcode.make(verification_url(resp.project_id, resp.merkle_root)).save(qr_buf, format="PNG")
    qr_buf.seek(0)
    img = XLImage(qr_buf)
    img.width = img.height = 110
    ws.add_image(img, "H1")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_pdf(resp: BudgetValidationResponse) -> bytes:
    """Layer 1 in PDF: tabella validata, controlli di budget, CEP-ID e QR verso l'Auditor Portal."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def eur(v: float) -> str:
        return f"{v:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")

    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7.5, leading=9)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"QUANTO {resp.cep_id}", author="QUANTO")

    qr_img = io.BytesIO()
    qrcode.make(verification_url(resp.project_id, resp.merkle_root)).save(qr_img, format="PNG")
    qr_img.seek(0)
    head = Table([[
        [Paragraph("<b>QUANTO — Budget validato</b>", styles["Title"]),
         Paragraph(f"Progetto: <b>{resp.project_id}</b> · Bando: <b>{resp.bando_id}</b>", styles["BodyText"]),
         Paragraph(f"CEP-ID: <b>{resp.cep_id}</b> · Conformity Score: <b>{resp.conformity_score}/100</b> · Stato: <b>{resp.status}</b>", styles["BodyText"]),
         Paragraph(f"Merkle Root: <font face='Courier' size='7'>{resp.merkle_root}</font>", styles["BodyText"])],
        Image(qr_img, 28 * mm, 28 * mm),
    ]], colWidths=[225 * mm, 35 * mm])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))

    rows = [["ID", "Descrizione", "Cat.", "Stato", "Richiesto €", "Ammesso €", "€/h", "Hash riga (SHA-256)"]]
    for i in resp.items:
        rows.append([i.item_id, Paragraph(i.description, small), i.category.value, i.status.value, eur(i.original_cost_eur), eur(i.computed_cost_eur),
                     eur(i.hourly_rate_computed) if i.hourly_rate_computed else "—", Paragraph(f"<font face='Courier'>{i.item_hash_sha256[:32]}…</font>", small)])
    rows.append(["", "", "", "TOTALE", eur(resp.total_requested_eur), eur(resp.total_approved_eur), "", ""])
    table = Table(rows, repeatRows=1, colWidths=[20 * mm, 62 * mm, 26 * mm, 40 * mm, 26 * mm, 26 * mm, 16 * mm, 50 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (4, 1), (6, -1), "RIGHT"), ("FONTNAME", (3, -1), (5, -1), "Helvetica-Bold"),
    ]))
    story = [head, Spacer(1, 6 * mm), table]

    evaluated = [c for c in resp.budget_checks if c.status != "NOT_EVALUATED"]
    if evaluated:
        story += [Spacer(1, 5 * mm), Paragraph("<b>Controlli sull'intero budget</b>", styles["Heading4"])]
        story += [Paragraph(f"[{c.status}] criterio {c.criterion} — {c.message}", small) for c in evaluated]
    story += [Spacer(1, 5 * mm), Paragraph(
        "Verifica ex-post: inquadrare il QR (Auditor Portal) o inserire il CEP-ID. "
        "La verifica ricalcola l'Albero di Merkle e lo confronta con la registrazione firmata.", small)]
    doc.build(story)
    return buf.getvalue()
