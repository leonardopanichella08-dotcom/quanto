"""Lettura dei fogli Excel dell'azienda: organico, libro cespiti, registro delle fatture, registro degli aiuti, budget e piano degli investimenti e delle assunzioni.

Un foglio di calcolo non ha uno schema fisso come un bilancio: si riconosce dalle intestazioni delle colonne. Ogni foglio riconosciuto produce **fatti** con il loro
tipo (``staff_member``, ``asset_line``, ``supplier_spend``…), ognuno con la fonte (foglio e riga). Un foglio che non si riconosce resta archiviato e basta:
non si indovina mai cosa contiene.

Regole:
- i **dati personali** (nome, cognome, codice fiscale) non entrano mai nei fatti: resta il token (stesso token delle buste paga della stessa persona);
- si leggono i **valori** delle celle, non le formule (un totale scritto come formula si ricalcola, non si copia);
- le righe dei **costi** si classificano con lo stesso criterio dei bilanci (categoria e voce di spesa): una riga che non si riconosce resta «da classificare»,
  non finisce in una categoria a caso.
"""
from __future__ import annotations

import io
import json
import re
from datetime import date, datetime
from typing import Any, Dict, Iterable, List, Optional, Tuple

from app.core import expense_types
from app.core.anonymizer import PrivacyAnonymizer
from app.core.fonte_c import parsers

Field = parsers.Field
CONF = 0.97                                 # celle di un foglio strutturato: nessuna ambiguità di lettura (come per un testo PDF nativo)
FACT_KEYS = ("staff_member", "asset_line", "dep_plan_line", "supplier_spend", "aid_received", "capex_plan", "hiring_plan", "budget_line", "trial_balance_line", "durc", "lease", "rd_credit", "loan",
             "rna_visura", "rna_aid", "shareholder", "de_minimis_declared", "tax_return", "bank_balance", "quote", "xbrl_balance")
_MONTHS = {"GEN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAG": 5, "GIU": 6, "LUG": 7, "AGO": 8, "SET": 9, "OTT": 10, "NOV": 11, "DIC": 12}


def _num(v: Any) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and not v.startswith("="):
        try:
            return float(v.replace(".", "").replace(",", ".")) if "," in v else float(v)
        except ValueError:
            return None
    return None


def _txt(v: Any) -> str:
    return "" if v is None or (isinstance(v, str) and v.startswith("=")) else str(v).strip()


def _year(v: Any) -> Optional[int]:
    if isinstance(v, (date, datetime)):
        return v.year
    m = re.match(r"(20\d{2})", str(v or "")) or re.match(r"\d{1,2}[/.\-]\d{1,2}[/.\-](20\d{2})", str(v or ""))
    return int(m.group(1)) if m else None


def _date(v: Any) -> Optional[str]:
    """Una data come testo ISO: da una cella data, da «2024-04-12» o da «12/04/2024»."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    t = str(v or "").strip()
    m = re.match(r"(20\d{2})-(\d{2})-(\d{2})", t)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = re.match(r"(\d{1,2})[/.\-](\d{1,2})[/.\-](20\d{2})", t)
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else None


def _table(ws, max_header_row: int = 8) -> Tuple[Optional[List[str]], List[Tuple[int, Dict[str, Any]]]]:
    """Intestazioni normalizzate (maiuscole, senza accenti) e righe come dizionari, con il numero di riga del foglio."""
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows[:max_header_row]):
        cells = [parsers.norm(_txt(c)) for c in r]
        if sum(1 for c in cells if c) >= 3 and sum(1 for c in r if isinstance(c, str) and c) >= 3:
            heads = cells
            out = []
            for j, row in enumerate(rows[i + 1:], i + 2):
                if not any(c is not None and c != "" for c in row):
                    continue
                out.append((j, {h: row[k] for k, h in enumerate(heads) if h and k < len(row)}))
            return heads, out
    return None, []


def _has(heads: Iterable[str], *needles: str) -> bool:
    hs = list(heads)
    return all(any(n in h for h in hs) for n in needles)


def _col(row: Dict[str, Any], *needles: str) -> Any:
    for h, v in row.items():
        if all(n in h for n in needles):
            return v
    return None


def _cost_class(description: str) -> Tuple[Optional[str], Optional[str]]:
    """Categoria e voce di spesa di una riga di costo, con lo stesso criterio dei bilanci; (None, None) se non si riconosce."""
    cat, _ = parsers.classify_cost(description)
    if cat is None:
        code = expense_types.classify_type(description)
        cat = expense_types.CATEGORY_OF.get(code) if code else None
    if cat is None:
        return None, None
    return cat, expense_types.classify_type(description, cat)


def _fact(key: str, payload: Dict[str, Any], sheet_no: int, snippet: str, pii: bool = False) -> Field:
    return Field(key, json.dumps(payload, ensure_ascii=False, sort_keys=True), CONF, sheet_no, snippet[:200], pii)


# ------------------------------------------------------------------------------------------------ i singoli fogli
def _staff(rows, no, name) -> List[Field]:
    out = []
    for r, row in rows:
        mat = _txt(_col(row, "MATRICOLA"))
        ral = _num(_col(row, "RAL"))
        level = _txt(_col(row, "LIVELLO"))
        if not mat or ral is None or not level:
            continue
        cf = _txt(_col(row, "CODICE FISCALE"))
        token = PrivacyAnonymizer.generate_opaque_token(cf) if cf else PrivacyAnonymizer.generate_opaque_token(f"{name}:{mat}")
        hired = _col(row, "ASSUNZIONE")
        out.append(_fact("staff_member", {
            "token": token, "ref": mat, "role": _txt(_col(row, "RUOLO")), "level": level, "contract": _txt(_col(row, "TIPO CONTRATTO")), "ccnl": _txt(_col(row, "CCNL APPLICATO") or _col(row, "CONTRATTO COLLETTIVO")) or None,
            "hours_week": _num(_col(row, "ORE SETTIMANALI")), "payments": _num(_col(row, "MENSILIT")), "ral_eur": round(ral, 2),
            "hired": hired.date().isoformat() if isinstance(hired, (date, datetime)) else None,
            "birth_year": _year(_col(row, "NASCITA")), "left": None}, no, f"{name}, riga {r}", pii=False))
    return out


def _assets(rows, no, name) -> List[Field]:
    out = []
    for r, row in rows:
        cost = _num(_col(row, "COSTO STORICO"))
        code = _txt(_col(row, "CODICE BENE"))
        if cost is None or not code:
            continue
        acc = _num(_col(row, "FONDO")) or 0.0
        desc = _txt(_col(row, "DESCRIZIONE"))
        out.append(_fact("asset_line", {"class": _txt(_col(row, "CLASSE")), "code": code, "description": desc, "supplier": _txt(_col(row, "FORNITORE")) or None,
                                        "purchase_year": _year(_col(row, "DATA ACQUISTO")), "rate": _num(_col(row, "ALIQUOTA")), "cost_eur": round(cost, 2), "accumulated_eur": round(acc, 2),
                                        "net_eur": round(cost - acc, 2), "expense_type": expense_types.classify_type(f"{_txt(_col(row, 'CLASSE'))} {desc}", "CAPITAL_ASSETS")}, no, f"{name}, riga {r}"))
    return out


def _dep_plan(rows, no, name) -> List[Field]:
    out = []
    for r, row in rows:
        desc = _txt(_col(row, "BENE"))
        years = {h: _num(v) for h, v in row.items() if re.fullmatch(r"20\d{2}", h or "")}
        if not desc or not years:
            continue
        out.append(_fact("dep_plan_line", {"class": _txt(_col(row, "CLASSE")), "description": desc, "net_eur": _num(_col(row, "VALORE NETTO")),
                                           "plan": {y: round(v or 0.0, 2) for y, v in years.items()}}, no, f"{name}, riga {r}"))
    return out


def _supplier_spend(rows, no, name) -> List[Field]:
    """Le fatture ricevute, sommate per anno e per voce di spesa (l'imponibile): migliaia di righe diventano poche decine di fatti, ognuno con il numero di fatture."""
    agg: Dict[Tuple[int, Optional[str], Optional[str]], List[float]] = {}
    for r, row in rows:
        amount = _num(_col(row, "IMPONIBILE"))
        y = _year(_col(row, "DATA"))
        if amount is None or y is None:
            continue
        cat, typ = _cost_class(_txt(_col(row, "DESCRIZIONE")))
        k = (y, typ, cat)
        a = agg.setdefault(k, [0.0, 0])
        a[0] += amount
        a[1] += 1
    return [_fact("supplier_spend", {"year": y, "expense_type": t, "category": c, "amount_eur": round(a[0], 2), "invoices": a[1]}, no, f"{name}: {a[1]} fatture")
            for (y, t, c), a in sorted(agg.items(), key=lambda kv: (kv[0][0], -kv[1][0]))]


def _aids(rows, no, name) -> List[Field]:
    out = []
    for r, row in rows:
        amount = _num(_col(row, "IMPORTO"))
        granted = _date(_col(row, "DATA CONCESSIONE"))
        if amount is None or granted is None:
            continue
        out.append(_fact("aid_received", {"granted": granted, "entity": _txt(_col(row, "ENTE")), "measure": _txt(_col(row, "MISURA")),
                                          "regime": _txt(_col(row, "REGIME")), "amount_eur": round(amount, 2), "cor": _txt(_col(row, "COR")) or None, "eligible_costs": _txt(_col(row, "SPESE AGEVOLATE")) or None}, no, f"{name}, riga {r}"))
    return out


def _capex(rows, heads, no, name) -> List[Field]:
    months = {h: (_MONTHS.get(h[:3]), _year(h[-4:])) for h in heads if re.fullmatch(r"[A-Z]{3} 20\d{2}", h or "")}
    out = []
    for r, row in rows:
        desc = _txt(_col(row, "INVESTIMENTO"))
        if not desc or parsers.norm(desc) == "TOTALE":
            continue
        per_year: Dict[int, float] = {}
        for h, (m, y) in months.items():
            v = _num(row.get(h))
            if v and y:
                per_year[y] = per_year.get(y, 0.0) + v
        for y, v in sorted(per_year.items()):
            out.append(_fact("capex_plan", {"description": desc, "year": y, "amount_eur": round(v, 2), "category": "CAPITAL_ASSETS", "expense_type": expense_types.classify_type(desc, "CAPITAL_ASSETS")}, no, f"{name}, riga {r}"))
    return out


def _hiring(rows, no, name) -> List[Field]:
    out = []
    for r, row in rows:
        cost = _num(_col(row, "COSTO AZIENDA"))
        start = _txt(_col(row, "MESE DI INGRESSO"))
        if cost is None or not start:
            continue
        out.append(_fact("hiring_plan", {"position": _txt(_col(row, "POSIZIONE")), "role": _txt(_col(row, "RUOLO")), "area": _txt(_col(row, "AREA")), "start": start[:7],
                                         "level": _txt(_col(row, "LIVELLO")), "annual_cost_eur": round(cost, 2)}, no, f"{name}, riga {r}"))
    return out


def _budget(ws, rows, no, name) -> List[Field]:
    title = _txt(ws.cell(row=1, column=1).value)
    m = re.search(r"20\d{2}", title)
    y = int(m.group(0)) if m else None
    out = []
    for r, row in rows:
        code, desc = _txt(_col(row, "CODICE")), _txt(_col(row, "CONTO"))
        months = [_num(v) for h, v in row.items() if h.startswith("BUDGET ") and h[7:10] in _MONTHS]          # il totale annuo è una formula: si sommano i dodici mesi
        total = sum(m for m in months if m is not None) if len([m for m in months if m is not None]) >= 12 else None
        if not code.startswith("3") or total is None or not desc or y is None:
            continue
        cat, typ = _cost_class(desc)
        out.append(_fact("budget_line", {"year": y, "code": code, "description": desc, "category": cat, "expense_type": typ, "amount_eur": round(total, 2)}, no, f"{name}, riga {r}"))
    return out


def _trial_balance(rows, no, name, year: Optional[int]) -> List[Field]:
    """Il bilancio di verifica del gestionale: i conti di costo dell'esercizio. Un conto economico parte da saldo zero a inizio anno (un conto patrimoniale no), e un costo ha saldo finale positivo in dare:
    sono i due segni che lo distinguono, qualunque sia il piano dei conti del gestionale. Un conto che non si riconosce come costo non entra."""
    if year is None:
        return []
    out = []
    for r, row in rows:
        desc = _txt(_col(row, "CONTO") or _col(row, "DESCRIZIONE"))
        final = _num(_col(row, "SALDO FINALE"))
        initial = _num(_col(row, "SALDO INIZIALE"))
        if not desc or final is None or final <= 0 or (initial is not None and abs(initial) > 0.005):
            continue
        if any(k in parsers.norm(desc) for k in ("INTERESS", "ONERI FINANZ", "IMPOSTE SUL REDDITO", "IRES", "IRAP", "PERDITE SU CAMBI")):          # oneri finanziari e imposte sul reddito non sono costi della produzione
            continue
        cat, typ = _cost_class(desc)
        if cat is None:
            continue
        out.append(_fact("trial_balance_line", {"year": year, "code": _txt(_col(row, "CODICE")), "description": desc, "category": cat, "expense_type": typ, "amount_eur": round(final, 2)}, no, f"{name}, riga {r}"))
    return out


# ------------------------------------------------------------------------------------------------ riconoscimento di una tabella
def _recognize(title: str, heads: List[str], rows, no: int, ws=None) -> List[Field]:
    """Cosa dice una tabella, dalle intestazioni delle sue colonne (vale per un foglio Excel e per un CSV esportato da un gestionale)."""
    if _has(heads, "MATRICOLA", "LIVELLO", "RAL"):
        return _staff(rows, no, title)
    if _has(heads, "CODICE BENE", "COSTO STORICO"):
        return _assets(rows, no, title)
    if "AMMORTAMENTO" in parsers.norm(title) and _has(heads, "BENE", "VALORE NETTO"):
        return _dep_plan(rows, no, title)
    if _has(heads, "PROTOCOLLO", "FORNITORE", "IMPONIBILE"):
        return _supplier_spend(rows, no, title)
    if _has(heads, "ENTE CONCEDENTE", "IMPORTO", "REGIME"):
        return _aids(rows, no, title)
    if heads and heads[0].startswith("INVESTIMENTO"):
        return _capex(rows, heads, no, title)
    if _has(heads, "POSIZIONE", "MESE DI INGRESSO"):
        return _hiring(rows, no, title)
    if ws is not None and _has(heads, "CODICE", "CONTO", "BUDGET ANNO"):
        return _budget(ws, rows, no, title)
    if _has(heads, "CODICE", "SALDO FINALE") and (_has(heads, "CONTO") or _has(heads, "DESCRIZIONE")):
        m = re.search(r"20\d{2}", title)
        return _trial_balance(rows, no, title, int(m.group(0)) if m else None)
    return []


# ------------------------------------------------------------------------------------------------ cartella di lavoro e CSV
def read_workbook(data: bytes) -> Tuple[str, List[Field], List[str]]:
    """(descrizione di ciò che contiene, fatti letti, fogli riconosciuti). Se nessun foglio si riconosce: ('', [], [])."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    except Exception:
        return "", [], []
    fields: List[Field] = []
    seen: List[str] = []
    for no, ws in enumerate(wb.worksheets, 1):
        heads, rows = _table(ws)
        if not heads or not rows:
            continue
        got = _recognize(ws.title, heads, rows, no, ws)
        if got:
            fields += got
            seen.append(ws.title)
    return ("Fogli letti: " + ", ".join(seen)) if fields else "", fields, seen


def read_csv(data: bytes, filename: str = "") -> Tuple[str, List[Field], List[str]]:
    """Un CSV esportato da un gestionale (bilancio di verifica, registro delle fatture…): stesse tabelle degli Excel, stessa lettura."""
    import csv
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            text = ""
    if not text or text.count("\n") < 2:
        return "", [], []
    head = text.splitlines()[0]
    delim = max((";", ",", "\t"), key=head.count)
    reader = list(csv.reader(text.splitlines(), delimiter=delim))
    if not reader or len(reader[0]) < 3:
        return "", [], []
    heads = [parsers.norm(h) for h in reader[0]]
    rows = []
    for j, cells in enumerate(reader[1:], 2):
        if not any(c.strip() for c in cells):
            continue
        rows.append((j, {h: (cells[k].strip() if k < len(cells) else "") for k, h in enumerate(heads) if h}))
    title = filename.rsplit(".", 1)[0] or "CSV"
    got = _recognize(title, heads, rows, 1, None)
    return (f"CSV letto: {filename}" if got else ""), got, ([title] if got else [])
