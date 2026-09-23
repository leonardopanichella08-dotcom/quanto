"""Caricamento, verifica e pubblicazione delle tabelle di Fonte B (Modulo 9.2).

Flusso: file (CSV o Excel) → lettura rigorosa riga per riga → insieme in BOZZA (con il file originale conservato) →
un manager lo controlla e lo PUBBLICA attestando la fonte ufficiale. Alla pubblicazione l'insieme precedente dello
stesso tipo e codice viene chiuso il giorno prima e resta consultabile per i calcoli retroattivi.

Le percentuali si scrivono «30%» oppure «0,30». Un numero come «30» senza il segno % è ambiguo e viene rifiutato:
meglio un errore visibile che un oneri sociali del 3000%.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

from app.core import events
from app.core.db import Conn, connect
from app.core.fonte_b import KNOWN_PARAMS

KINDS = ("CCNL", "PARAMS", "AMORTIZATION", "BENCHMARK")
KIND_LABEL = {"CCNL": "Costo del lavoro per CCNL e livello", "PARAMS": "Parametri (tempo determinato, occasionali…)",
              "AMORTIZATION": "Aliquote d'ammortamento per categoria di bene", "BENCHMARK": "Benchmark di prezzo e di tariffa"}

# colonne accettate (nome canonico → sinonimi)
COLUMNS: Dict[str, Dict[str, Tuple[str, ...]]] = {
    "CCNL": {"level": ("level", "livello", "inquadramento"), "standard_hours": ("standard_hours", "ore", "ore_annue", "ore_lavorabili"),
             "social_charges_pct": ("social_charges_pct", "oneri", "oneri_sociali", "oneri_pct"), "tfr_pct": ("tfr_pct", "tfr")},
    "PARAMS": {"param_key": ("param_key", "parametro", "chiave"), "value": ("value", "valore"), "unit": ("unit", "unita", "unità")},
    "AMORTIZATION": {"category_code": ("category_code", "categoria", "codice"), "description": ("description", "descrizione"),
                     "rate_pct": ("rate_pct", "aliquota", "aliquota_pct")},
    "BENCHMARK": {"bench_kind": ("bench_kind", "tipo"), "category_code": ("category_code", "categoria", "codice"),
                  "description": ("description", "descrizione"), "reference_eur": ("reference_eur", "valore", "valore_eur", "riferimento")},
}
REQUIRED = {"CCNL": ["level", "standard_hours", "social_charges_pct", "tfr_pct"], "PARAMS": ["param_key", "value"],
            "AMORTIZATION": ["category_code", "rate_pct"], "BENCHMARK": ["bench_kind", "category_code", "reference_eur"]}


class FonteBFileError(ValueError):
    """Il file non è valido: ``errors`` elenca le righe con il motivo."""

    def __init__(self, message: str, errors: Optional[List[Dict[str, Any]]] = None):
        super().__init__(message)
        self.errors = errors or []


def template_csv(kind: str) -> str:
    """Solo l'intestazione: nessun valore d'esempio (non si suggeriscono numeri che potrebbero finire nei calcoli)."""
    if kind not in KINDS:
        raise KeyError(kind)
    return ";".join(COLUMNS[kind]) + "\n"


# ------------------------------------------------------------------------------------------------ lettura file
def _norm(h: Any) -> str:
    return re.sub(r"[^a-z0-9_]+", "_", str(h or "").strip().lower()).strip("_")


def _read_table(filename: str, data: bytes) -> List[Dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = ";" if text.splitlines()[0].count(";") >= text.splitlines()[0].count(",") else ","
        rows = list(csv.reader(io.StringIO(text), delimiter=dialect))
    else:
        raise FonteBFileError("Formato non supportato: usa un file .csv o .xlsx")
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if len(rows) < 2:
        raise FonteBFileError("Il file non contiene righe di dati (serve l'intestazione e almeno una riga)")
    header = [_norm(h) for h in rows[0]]
    return [dict(zip(header, r)) | {"_line": i + 2} for i, r in enumerate(rows[1:])]


def _pick(row: Dict[str, Any], names: Tuple[str, ...]) -> Any:
    for n in names:
        if n in row and row[n] is not None and str(row[n]).strip() != "":
            return row[n]
    return None


def _decimal(value: Any, what: str, pct: bool = False, positive: bool = False) -> Decimal:
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        text = str(value)
    else:
        text = str(value).strip().replace("€", "").replace(" ", "")
    has_pct = text.endswith("%")
    text = text.rstrip("%")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        n = Decimal(text)
    except InvalidOperation:
        raise ValueError(f"{what}: «{value}» non è un numero") from None
    if pct:
        if has_pct:
            n = n / 100
        elif n > 1:
            raise ValueError(f"{what}: «{value}» è ambiguo. Scrivi {n}% oppure {n / 100}")
        if not 0 <= n <= 1:
            raise ValueError(f"{what}: {n} fuori intervallo (0-100%)")
    if positive and n <= 0:
        raise ValueError(f"{what}: deve essere maggiore di zero")
    return n


def parse_rows(kind: str, filename: str, data: bytes) -> List[Dict[str, Any]]:
    """Legge e valida il file. Restituisce righe pulite oppure solleva ``FonteBFileError`` con tutti gli errori."""
    if kind not in KINDS:
        raise FonteBFileError(f"Tipo sconosciuto: {kind}")
    raw = _read_table(filename, data)
    cols = COLUMNS[kind]
    missing = [c for c in REQUIRED[kind] if not any(any(n in r for n in cols[c]) for r in raw[:1])]
    if missing:
        raise FonteBFileError("Colonne mancanti: " + ", ".join(missing) + ". Attese: " + ", ".join(cols))
    out: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    seen = set()
    for r in raw:
        line = r["_line"]
        try:
            v = {c: _pick(r, names) for c, names in cols.items()}
            for c in REQUIRED[kind]:
                if v[c] is None:
                    raise ValueError(f"campo «{c}» vuoto")
            if kind == "CCNL":
                row = {"level": str(v["level"]).strip(), "standard_hours": int(_decimal(v["standard_hours"], "ore", positive=True)),
                       "social_charges_pct": _decimal(v["social_charges_pct"], "oneri", pct=True), "tfr_pct": _decimal(v["tfr_pct"], "tfr", pct=True)}
                key = row["level"]
            elif kind == "PARAMS":
                k = str(v["param_key"]).strip()
                if k not in KNOWN_PARAMS:
                    raise ValueError(f"parametro «{k}» sconosciuto. Noti: {', '.join(KNOWN_PARAMS)}")
                is_pct = k.endswith("_pct")
                row = {"param_key": k, "value": _decimal(v["value"], k, pct=is_pct, positive=True), "unit": (str(v["unit"]).strip() if v["unit"] else KNOWN_PARAMS[k][1])}
                key = k
            elif kind == "AMORTIZATION":
                row = {"category_code": str(v["category_code"]).strip().upper(), "description": (str(v["description"]).strip() if v["description"] else None),
                       "rate_pct": _decimal(v["rate_pct"], "aliquota", pct=True, positive=True)}
                key = row["category_code"]
            else:
                bk = str(v["bench_kind"]).strip().upper()
                if bk not in ("PRICE", "DAILY_RATE"):
                    raise ValueError("tipo deve essere PRICE oppure DAILY_RATE")
                row = {"bench_kind": bk, "category_code": str(v["category_code"]).strip().upper(),
                       "description": (str(v["description"]).strip() if v["description"] else None),
                       "reference_eur": _decimal(v["reference_eur"], "valore", positive=True)}
                key = (bk, row["category_code"])
            if key in seen:
                raise ValueError(f"riga duplicata per {key}")
            seen.add(key)
            out.append(row)
        except ValueError as exc:
            errors.append({"line": line, "error": str(exc)})
    if errors:
        raise FonteBFileError(f"{len(errors)} righe non valide: nulla è stato salvato", errors)
    return out


# ------------------------------------------------------------------------------------------------ bozze
def _dataset_row(conn: Conn, dataset_id: int):
    return conn.execute("SELECT * FROM fonte_b_datasets WHERE id=?", (dataset_id,)).fetchone()


def _public(r) -> Dict[str, Any]:
    d = {k: r[k] for k in r.keys() if k != "file_data"}
    for k in ("valid_from", "valid_to"):
        d[k] = d[k].isoformat() if d.get(k) else None
    d["has_file"] = bool(r["file_sha256"])
    return d


def create_draft(kind: str, code: str, name: str, version: str, valid_from: date, valid_to: Optional[date], source_name: str,
                 source_url: Optional[str], source_ref: Optional[str], filename: str, data: bytes, actor: str, notes: Optional[str] = None,
                 official: bool = False) -> Dict[str, Any]:
    code = code.strip().upper()
    if not code or not name.strip() or not version.strip() or not source_name.strip():
        raise FonteBFileError("Servono codice, nome, versione e nome della fonte")
    if kind == "CCNL" and not re.fullmatch(r"[A-Z0-9_]{2,40}", code):
        raise FonteBFileError("Il codice del CCNL usa solo lettere maiuscole, cifre e trattino basso (es. TERZO_SETTORE)")
    rows = parse_rows(kind, filename, data)
    with connect() as conn:
        if conn.execute("SELECT 1 FROM fonte_b_datasets WHERE kind=? AND code=? AND version=?", (kind, code, version.strip())).fetchone():
            raise FonteBFileError(f"Esiste già la versione «{version}» di {kind} {code}")
        did = conn.execute(
            "INSERT INTO fonte_b_datasets (kind, code, name, version, valid_from, valid_to, status, official, source_name, source_url, source_ref, "
            "file_name, file_sha256, file_data, notes, created_by, created_at) VALUES (?,?,?,?,?,?,'DRAFT',?,?,?,?,?,?,?,?,?,?) RETURNING id",
            (kind, code, name.strip(), version.strip(), valid_from, valid_to, official, source_name.strip(), source_url, source_ref, filename,
             hashlib.sha256(data).hexdigest(), data, notes, actor, events.now_iso())).fetchone()["id"]
        _insert_rows(conn, did, kind, rows)
        out = _public(_dataset_row(conn, did))
    out["rows"] = len(rows)
    return out


def _insert_rows(conn: Conn, did: int, kind: str, rows: List[Dict[str, Any]]) -> None:
    if kind == "CCNL":
        conn.executemany("INSERT INTO fonte_b_ccnl (dataset_id, level, standard_hours, social_charges_pct, tfr_pct) VALUES (?,?,?,?,?)",
                         [(did, r["level"], r["standard_hours"], r["social_charges_pct"], r["tfr_pct"]) for r in rows])
    elif kind == "PARAMS":
        conn.executemany("INSERT INTO fonte_b_params (dataset_id, param_key, value, unit) VALUES (?,?,?,?)", [(did, r["param_key"], r["value"], r["unit"]) for r in rows])
    elif kind == "AMORTIZATION":
        conn.executemany("INSERT INTO fonte_b_amort (dataset_id, category_code, description, rate_pct) VALUES (?,?,?,?)",
                         [(did, r["category_code"], r["description"], r["rate_pct"]) for r in rows])
    else:
        conn.executemany("INSERT INTO fonte_b_benchmarks (dataset_id, bench_kind, category_code, description, reference_eur) VALUES (?,?,?,?,?)",
                         [(did, r["bench_kind"], r["category_code"], r["description"], r["reference_eur"]) for r in rows])


ROW_TABLE = {"CCNL": "fonte_b_ccnl", "PARAMS": "fonte_b_params", "AMORTIZATION": "fonte_b_amort", "BENCHMARK": "fonte_b_benchmarks"}


def list_datasets() -> List[Dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute("SELECT * FROM fonte_b_datasets ORDER BY kind, code, valid_from DESC, id DESC").fetchall()
        out = []
        for r in rows:
            d = _public(r)
            d["rows"] = conn.execute(f"SELECT COUNT(*) c FROM {ROW_TABLE[r['kind']]} WHERE dataset_id=?", (r["id"],)).fetchone()["c"]
            out.append(d)
    return out


def dataset_detail(dataset_id: int) -> Optional[Dict[str, Any]]:
    with connect() as conn:
        r = _dataset_row(conn, dataset_id)
        if r is None:
            return None
        rows = [dict(x) for x in conn.execute(f"SELECT * FROM {ROW_TABLE[r['kind']]} WHERE dataset_id=? ORDER BY 2", (dataset_id,)).fetchall()]
    for x in rows:
        for k, v in list(x.items()):
            if isinstance(v, Decimal):
                x[k] = str(v)
    return {**_public(r), "data": rows}


def original_file(dataset_id: int) -> Optional[Dict[str, Any]]:
    with connect() as conn:
        r = _dataset_row(conn, dataset_id)
    if r is None or r["file_data"] is None:
        return None
    return {"name": r["file_name"], "data": bytes(r["file_data"])}


def delete_draft(dataset_id: int) -> bool:
    with connect() as conn:
        r = _dataset_row(conn, dataset_id)
        if r is None:
            return False
        if r["status"] != "DRAFT":
            raise PermissionError("Solo le bozze si eliminano: un insieme pubblicato resta per i calcoli retroattivi")
        conn.execute("DELETE FROM fonte_b_datasets WHERE id=?", (dataset_id,))
    return True


def publish(dataset_id: int, actor: str, attest_official: bool) -> Dict[str, Any]:
    """Pubblica una bozza. ``attest_official`` = il manager attesta che i valori corrispondono alla fonte indicata."""
    with connect() as conn:
        conn.lock(7_419_003)
        r = _dataset_row(conn, dataset_id)
        if r is None:
            raise KeyError(dataset_id)
        if r["status"] != "DRAFT":
            raise PermissionError("L'insieme è già pubblicato")
        if not attest_official:
            raise PermissionError("Per pubblicare devi attestare che i valori sono quelli della fonte ufficiale indicata")
        if not (r["source_url"] or r["source_ref"]):
            raise PermissionError("Indica l'indirizzo o il riferimento (numero, data, articolo) del documento ufficiale da cui vengono i valori")
        prev = conn.execute(
            "SELECT id, valid_from, valid_to FROM fonte_b_datasets WHERE kind=? AND code=? AND status IN ('PUBLISHED','SUPERSEDED') AND id<>? ORDER BY valid_from",
            (r["kind"], r["code"], dataset_id)).fetchall()
        for p in prev:
            p_end = p["valid_to"] or date.max
            r_end = r["valid_to"] or date.max
            if p["valid_from"] <= r_end and r["valid_from"] <= p_end:
                # sovrapposizione: si accetta solo se la precedente è aperta e parte prima → si chiude il giorno prima
                if p["valid_to"] is None and p["valid_from"] < r["valid_from"]:
                    conn.execute("UPDATE fonte_b_datasets SET valid_to=?, status='SUPERSEDED' WHERE id=?", (r["valid_from"] - timedelta(days=1), p["id"]))
                else:
                    raise PermissionError(f"Il periodo di validità si sovrappone all'insieme {p['id']} già pubblicato")
        conn.execute("UPDATE fonte_b_datasets SET status='PUBLISHED', official=TRUE, published_by=?, published_at=? WHERE id=?", (actor, events.now_iso(), dataset_id))
        return _public(_dataset_row(conn, dataset_id))
