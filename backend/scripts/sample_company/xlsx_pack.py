"""Fogli Excel dell'azienda di prova: decine o centinaia di righe e colonne, con formule e totali, come quelli di un cliente vero."""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path
from typing import Dict, List

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import assets as AS
from . import bp
from . import ledger as LG
from . import model as M
from . import people as P
from .xl import DATE, INT, MONEY, PCT, add_sheet, euro, new_book

CO = M.COMPANY
MESI = ["Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno", "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]
MS = ["Gen", "Feb", "Mar", "Apr", "Mag", "Giu", "Lug", "Ago", "Set", "Ott", "Nov", "Dic"]


# ================================================================================================ registro aiuti
def aiuti(path: Path) -> Path:
    wb = new_book()
    head = ["Ente concedente", "Misura", "Base normativa", "Data concessione", "Data incasso/fruizione", "Importo (EUR)", "Regime", "Codice COR (prova)", "Impresa unica", "Spese agevolate", "Note"]
    rows = [
        ["Camera di Commercio di Torino", "Voucher digitalizzazione PMI", "Reg. UE 2023/2831", date(2024, 4, 12), date(2024, 9, 30), 10000, "de minimis", "COR-0000001", CO["short"], "Consulenza e licenze 2024", "Rendicontazione chiusa il 18/09/2024"],
        ["Regione Piemonte", "Contributo partecipazione fiere", "Reg. UE 2023/2831", date(2025, 10, 3), date(2025, 12, 15), 6500, "de minimis", "COR-0000002", CO["short"], "Fiera di settore a Milano, 2025", "Rendicontazione chiusa il 28/11/2025"],
        ["Agenzia delle Entrate", "Credito d'imposta ricerca e sviluppo 2022", "Art. 1 c. 198 L. 160/2019", date(2023, 6, 30), date(2023, 7, 17), M.GRANTS[0], "credito d'imposta (esenzione)", "-", CO["short"], "Spese di ricerca e sviluppo 2022", "Compensazione in F24 codice 6938"],
        ["Agenzia delle Entrate", "Credito d'imposta ricerca e sviluppo 2023", "Art. 1 c. 198 L. 160/2019", date(2024, 6, 30), date(2024, 7, 16), M.GRANTS[1], "credito d'imposta (esenzione)", "-", CO["short"], "Spese di ricerca e sviluppo 2023", "Compensazione in F24 codice 6938"],
        ["Agenzia delle Entrate", "Credito d'imposta ricerca e sviluppo 2024", "Art. 1 c. 198 L. 160/2019", date(2025, 6, 30), date(2025, 7, 16), M.GRANTS[2], "credito d'imposta (esenzione)", "-", CO["short"], "Spese di ricerca e sviluppo 2024", "Compensazione in F24 codice 6938"],
        ["Agenzia delle Entrate", "Credito d'imposta ricerca e sviluppo 2025", "Art. 1 c. 198 L. 160/2019", date(2026, 6, 30), "", M.GRANTS[3], "credito d'imposta (esenzione)", "-", CO["short"], "Spese di ricerca e sviluppo 2025", "Da utilizzare in tre quote annuali da luglio 2026"],
        ["MIMIT", "Nuova Sabatini: contributo in conto impianti", "D.L. 69/2013", date(2021, 5, 20), date(2021, 12, 10), 4900, "esenzione (Reg. UE 651/2014)", "-", CO["short"], "Acquisto di server e apparati", "Contributo erogato in un'unica soluzione"],
        ["Ministero del Lavoro", "Esonero contributivo per assunzione di giovani under 36", "L. 178/2020", date(2022, 2, 1), date(2022, 2, 28), 3200, "esenzione (Reg. UE 651/2014)", "-", CO["short"], "Assunzioni 2022", "Fruito in compensazione con il DM10"],
        ["Fondimpresa", "Conto formazione: finanziamento corsi", "Accordo interconfederale", date(2023, 9, 14), date(2023, 12, 20), 5200, "de minimis", "COR-0000003", CO["short"], "Formazione tecnica 2023", "Concesso nel 2023; rientra nel triennio 2023-2025"],
        ["INAIL", "Riduzione del premio per prevenzione (OT23)", "D.M. 3 dicembre 2021", date(2025, 3, 31), date(2025, 3, 31), 1850, "riduzione del premio", "-", CO["short"], "Interventi di prevenzione 2024", "Riduzione applicata in autoliquidazione"],
    ]
    # il de minimis del 2023 (Fondimpresa) non compare nella dichiarazione del triennio 2024-2026 per la finestra: coerente con l'anno di concessione
    rows[8][5] = 5200
    add_sheet(wb, "Registro aiuti", head, rows, [30, 44, 26, 16, 18, 14, 28, 18, 30, 34, 44], {5: MONEY}, note="Registro degli aiuti di Stato e delle agevolazioni ricevute (aggiornato al 30/09/2026). Il de minimis è soggetto al massimale di 300.000 EUR su tre esercizi.", freeze="A4")
    ws = wb["Registro aiuti"]
    ws.cell(row=len(rows) + 5, column=1, value="Aiuti de minimis nel triennio 2024-2026").font = Font(bold=True)
    ws.cell(row=len(rows) + 5, column=6, value='=SUMIFS(F4:F13,G4:G13,"de minimis",D4:D13,">="&DATE(2024,1,1))').number_format = MONEY
    ws.cell(row=len(rows) + 6, column=1, value="Margine residuo rispetto al massimale").font = Font(bold=True)
    ws.cell(row=len(rows) + 6, column=6, value=f"=300000-F{len(rows) + 5}").number_format = MONEY
    # domande presentate
    apps = [["2019-03-04", "Regione Piemonte", "Bando Voucher digitali", 8000, "Non ammessa", "Graduatoria: fondi esauriti"], ["2020-05-18", "MIMIT", "Nuova Sabatini", 7000, "Concessa", "Contributo erogato nel 2021"],
            ["2021-02-22", "Camera di Commercio di Torino", "Voucher digitalizzazione 2021", 6000, "Non ammessa", "Domanda tardiva"], ["2022-03-10", "Camera di Commercio di Torino", "Voucher formazione 4.0", 4000, "Ritirata", "Cambio di progetto"],
            ["2023-04-27", "Fondimpresa", "Avviso formazione continua", 5200, "Concessa", "Registrata come de minimis"], ["2024-02-08", "Camera di Commercio di Torino", "Voucher digitalizzazione 2024", 10000, "Concessa", "COR-0000001"],
            ["2024-09-30", "Regione Piemonte", "Bando innovazione PMI", 90000, "Non ammessa", "Punteggio inferiore alla soglia"], ["2025-05-14", "Regione Piemonte", "Contributo partecipazione fiere 2025", 6500, "Concessa", "COR-0000002"],
            ["2026-03-10", "Camera di Commercio di Torino", "Voucher digitalizzazione 2026", 12000, "In istruttoria", "Presentata via PEC"], ["2026-05-22", "Regione Piemonte", "Bando SWIch 2026", 140000, "Non presentata", "Verifica dei requisiti in corso"],
            ["2026-07-01", "Invitalia", "Smart&Start Italia", 0, "Non ammissibile", "L'impresa ha più di 5 anni dall'iscrizione"], ["2026-09-15", "Ministero (bando innovazione)", "Progetto Atlas - candidatura", 143950, "In preparazione", "Bozza interna versione 0.7"]]
    add_sheet(wb, "Domande presentate", ["Data", "Ente", "Bando", "Importo richiesto (EUR)", "Esito", "Note"], [[date.fromisoformat(a[0])] + a[1:] for a in apps], [14, 34, 38, 22, 18, 44], {3: MONEY})
    # calcolo triennio mobile
    dm = [(date(2023, 9, 14), 5200), (date(2024, 4, 12), 10000), (date(2025, 10, 3), 6500)]
    calc = []
    d0 = date(2023, 1, 1)
    for k in range(48):
        mm = (d0.month - 1 + k) % 12 + 1
        yy = d0.year + (d0.month - 1 + k) // 12
        ref = date(yy, mm, 1)
        calc.append([ref, f'=SUMIFS(\'Registro aiuti\'!$F$4:$F$13,\'Registro aiuti\'!$G$4:$G$13,"de minimis",\'Registro aiuti\'!$D$4:$D$13,">"&EDATE(A{k + 2},-36),\'Registro aiuti\'!$D$4:$D$13,"<="&A{k + 2})', f"=300000-B{k + 2}"])
    add_sheet(wb, "Triennio mobile", ["Data di riferimento (1° del mese)", "Aiuti de minimis nei 36 mesi precedenti (EUR)", "Margine residuo (EUR)"], calc, [34, 44, 26], {0: DATE, 1: MONEY, 2: MONEY})
    # impresa unica e credito R&S
    add_sheet(wb, "Impresa unica", ["Verifica", "Esito", "Fonte"], [["Un socio detiene la maggioranza dei diritti di voto di un'altra impresa", "No", "Visura camerale"], ["Diritto di nominare la maggioranza degli amministratori di un'altra impresa", "No", "Visura camerale"],
                                                                        ["Contratti o clausole statutarie di influenza dominante", "No", "Statuto"], ["Partecipazioni della società in altre imprese", "Nessuna", "Bilancio, nota integrativa"], ["Operazioni straordinarie nel triennio", "Nessuna", "Verbali"]], [70, 14, 30])
    cr = []
    for i_, y in enumerate((2022, 2023, 2024, 2025)):
        base = M.GRANTS[i_] * (5 if y == 2022 else 10)
        cr.append([y, base, 0.2 if y == 2022 else 0.1, f"=B{i_ + 2}*C{i_ + 2}", "Certificata" if y < 2025 else "Certificazione in corso"])
    add_sheet(wb, "Credito R&S", ["Esercizio", "Costi ammissibili (EUR)", "Aliquota", "Credito maturato (EUR)", "Stato"], cr, [12, 26, 12, 24, 26], {1: MONEY, 2: PCT, 3: MONEY})
    wb.save(path)
    return path


# ================================================================================================ organico
def _monthly_gross(y: int) -> Dict[str, List[int]]:
    """Retribuzioni lorde mensili 2025 (centesimi) per dipendente: le somme tornano al conto «Salari e stipendi» del bilancio."""
    weights = LG.PROFILES["payroll"]
    cells = []
    keys = []
    for e in P.EMPLOYEES:
        hired = date.fromisoformat(e["hired"])
        for m in range(12):
            first = date(y, m + 1, 1)
            if hired > date(y, m + 1, 28):
                w = 0.0
            elif hired > first:
                w = (30 - hired.day) / 30
            else:
                w = 1.0
            cells.append(w * e["cost"] * weights[m])
            keys.append((e["id"], m))
    total = M.acct("Salari e stipendi", y) * 100
    parts = LG.alloc(total, cells)
    out: Dict[str, List[int]] = {e["id"]: [0] * 12 for e in P.EMPLOYEES}
    for (eid, m), v in zip(keys, parts):
        out[eid][m] = v
    return out


def organico(path: Path) -> Path:
    wb = new_book()
    head = ["Matricola", "Cognome", "Nome", "Codice fiscale", "Sesso", "Data di nascita", "Ruolo", "Livello CCNL", "Tipo contratto", "Data assunzione", "Ore settimanali", "Mensilità", "RAL (EUR)", "Retribuzione mensile lorda (EUR)",
            "Costo azienda annuo stimato (EUR)", "Centro di costo", "Sede", "Note"]
    rows = []
    for e in P.EMPLOYEES:
        ral = round(e["cost"] / 1.38 / 100) * 100
        cc = "Sviluppo" if any(k in e["role"] for k in ("viluppatore", "analyst", "scientist", "UX")) else ("Commerciale" if "Commerciale" in e["role"] else "Direzione e staff")
        rows.append([e["id"], e["surname"], e["name"], e["cf"], "M" if e["male"] else "F", e["birth"], e["role"], e["level"], e["contract"], date.fromisoformat(e["hired"]), e["hours"], 14, ral, f"=M{len(rows) + 2}/L{len(rows) + 2}", e["cost"], cc,
                     "Via Garibaldi 28" if e["id"] < "M008" else "Via Garibaldi 28 / Corso Francia 114", "Part time 75%" if e["hours"] == 30 else ("Apprendistato professionalizzante" if e["contract"] == "Apprendistato" else "")])
    add_sheet(wb, "Organico", head, rows, [10, 14, 12, 20, 6, 14, 32, 10, 16, 14, 10, 10, 14, 18, 20, 18, 30, 30], {5: DATE, 9: DATE, 12: INT, 13: MONEY, 14: INT}, total_cols=[12, 14], total_label="Totale")
    gross = _monthly_gross(2025)
    rows = [[e["id"], f"{e['surname']} {e['name']}"] + [euro(v) for v in gross[e["id"]]] + [f"=SUM(C{k + 2}:N{k + 2})"] for k, e in enumerate(P.EMPLOYEES)]
    add_sheet(wb, "Retribuzioni lorde 2025", ["Matricola", "Dipendente"] + MESI + ["Totale anno"], rows, [10, 26] + [13] * 12 + [16], {k: MONEY for k in range(2, 15)}, total_cols=list(range(2, 15)), freeze="C2")
    # oneri, TFR e costo per dipendente
    wages = {e["id"]: sum(gross[e["id"]]) for e in P.EMPLOYEES}
    tot = sum(wages.values())
    soc = LG.alloc(M.acct("Oneri sociali", 2025) * 100, [wages[e["id"]] for e in P.EMPLOYEES])
    tfr = LG.alloc(M.acct("Trattamento di fine rapporto", 2025) * 100, [wages[e["id"]] for e in P.EMPLOYEES])
    oth = LG.alloc(M.acct("Altri costi del personale", 2025) * 100, [wages[e["id"]] for e in P.EMPLOYEES])
    rows = [[e["id"], f"{e['surname']} {e['name']}", euro(wages[e["id"]]), euro(soc[k]), euro(tfr[k]), euro(oth[k]), f"=SUM(C{k + 2}:F{k + 2})"] for k, e in enumerate(P.EMPLOYEES)]
    add_sheet(wb, "Costo del lavoro 2025", ["Matricola", "Dipendente", "Retribuzioni lorde", "Oneri sociali", "TFR", "Altri costi (buoni pasto, welfare)", "Costo totale"], rows, [10, 26, 18, 16, 14, 24, 16], {k: MONEY for k in range(2, 7)}, total_cols=[2, 3, 4, 5, 6])
    # cedolini 2026 gennaio-agosto
    from .docs import _irpef
    rows = []
    for k, e in enumerate(P.EMPLOYEES):
        base = {"M003": 2980.77, "M006": 2564.10}.get(e["id"], round(e["cost"] / 1.38 / 14 * 1.012, 2))
        hired = date.fromisoformat(e["hired"])
        row = [e["id"], f"{e['surname']} {e['name']}"]
        for m in range(8):
            if hired > date(2026, m + 1, 28):
                row.append(0)
            else:
                row.append(round(base * (2 if m == 5 else 1), 2))
        row.append(f"=SUM(C{k + 2}:J{k + 2})")
        inps, taxable, gt, dt, nt, reg, com = _irpef(base, 14)
        row.append(round(base - inps - nt - reg - com, 2))
        rows.append(row)
    add_sheet(wb, "Cedolini 2026 (lordo)", ["Matricola", "Dipendente"] + MESI[:8] + ["Totale gen-ago", "Netto mensile tipo (EUR)"], rows, [10, 26] + [13] * 8 + [16, 20], {k: MONEY for k in range(2, 11)}, total_cols=list(range(2, 10)), freeze="C2")
    # presenze
    r_ = LG.rng("presenze")
    rows = []
    for k, e in enumerate(P.EMPLOYEES):
        ferie = round(22 * (e["hours"] / 40), 1)
        rows.append([e["id"], f"{e['surname']} {e['name']}", 2080 if e["hours"] == 40 else 1560, ferie, round(ferie * r_.uniform(0.82, 1.0), 1), 56 if e["hours"] == 40 else 42, round(r_.uniform(8, 44), 1), r_.randint(0, 6), r_.randint(0, 24), r_.randint(0, 3)])
    add_sheet(wb, "Presenze 2025", ["Matricola", "Dipendente", "Ore contrattuali annue", "Ferie maturate (giorni)", "Ferie godute (giorni)", "Permessi ROL maturati (ore)", "Permessi goduti (ore)", "Giorni di malattia", "Ore di straordinario", "Giorni di smart working extra"],
              rows, [10, 26, 18, 18, 18, 22, 18, 16, 18, 24], {k: '#,##0.0' for k in (3, 4, 6)}, total_cols=[2, 3, 4, 5, 6, 7, 8, 9])
    ev = [[e["hired"], e["id"], "Assunzione", e["role"], f"Livello {e['level']}, {e['contract']}"] for e in P.EMPLOYEES if e["hired"] >= "2022-01-01"]
    ev += [["2024-12-31", "M009", "Trasformazione", "Conferma", "Conferma al termine del periodo di prova"], ["2025-06-30", "M-ex1", "Cessazione", "Sviluppatore", "Dimissioni volontarie, preavviso rispettato"], ["2024-03-31", "M-ex2", "Cessazione", "Impiegato amministrativo", "Dimissioni volontarie"]]
    add_sheet(wb, "Assunzioni e cessazioni", ["Data", "Matricola", "Evento", "Ruolo", "Note"], [[date.fromisoformat(a[0])] + a[1:] for a in sorted(ev)], [14, 12, 18, 30, 50])
    plan = bp.staff()
    rows = [[pr["id"], pr["role"], pr["area"], f"{2026 + pr['start'] // 12}-{pr['start'] % 12 + 1:02d}", pr["level"], pr["cost"]] for pr in plan["people"] if pr["new"]]
    add_sheet(wb, "Piano assunzioni 2026-28", ["Posizione", "Ruolo", "Area", "Mese di ingresso", "Livello", "Costo azienda annuo di partenza (EUR)"], rows, [12, 36, 28, 18, 10, 32], {5: INT}, total_cols=[5])
    wb.save(path)
    return path


# ================================================================================================ cespiti
def cespiti(path: Path) -> Path:
    wb = new_book()
    head = ["Classe", "Codice bene", "Descrizione", "Fornitore", "Data acquisto", "Aliquota", "Costo storico (EUR)", "Fondo ammortamento 31/12/2025 (EUR)", "Valore residuo (EUR)", "Ammortamento 2025 (EUR)"]
    rows = []
    for kind, tag in (("intang", "IMM"), ("tang", "TNG")):
        for n, r in enumerate([x for x in AS.register(kind)[2025] if x["cost_close"] > 0], 1):
            rows.append([r["class_name"], f"{tag}-{r['class'] + 1}{n:02d}", r["item"], r["supplier"], r["date"], r["rate"], r["cost_close"], r["fund_close"], f"=G{len(rows) + 2}-H{len(rows) + 2}", r["dep"]])
    add_sheet(wb, "Libro cespiti 2025", head, rows, [44, 12, 62, 42, 14, 10, 18, 26, 18, 20], {4: DATE, 5: PCT, 6: INT, 7: INT, 8: INT, 9: INT}, total_cols=[6, 7, 8, 9], total_label="Totale immobilizzazioni (immateriali e materiali)")
    for y in (2022, 2023, 2024, 2025):
        rr = []
        for kind in ("intang", "tang"):
            for r in [x for x in AS.register(kind)[y] if x["cost_close"] > 0]:
                rr.append([r["class_name"], r["item"], r["date"], r["rate"], r["cost_open"], r["capex"], r["cost_close"], r["fund_open"], r["dep"], r["fund_close"], r["net_close"]])
        add_sheet(wb, f"Movimenti {y}", ["Classe", "Bene", "Acquisto", "Aliquota", "Costo 1/1", "Incrementi", "Costo 31/12", "Fondo 1/1", "Ammortamento", "Fondo 31/12", "Valore netto"], rr, [44, 62, 14, 10, 14, 14, 14, 14, 14, 14, 14],
                  {2: DATE, 3: PCT, **{k: INT for k in range(4, 11)}}, total_cols=list(range(4, 11)))
    rows = []
    for kind in ("intang", "tang"):
        for r in [x for x in AS.projected_depreciation(kind) if x["net_close"] > 0]:
            rows.append([r["class_name"], r["item"], r["net_close"]] + [r["plan"][y] for y in (2026, 2027, 2028, 2029, 2030)] + [r["after"]])
    add_sheet(wb, "Piano ammortamento 2026-30", ["Classe", "Bene", "Valore netto 31/12/2025", "2026", "2027", "2028", "2029", "2030", "Residuo dopo il 2030"], rows, [44, 62, 18, 12, 12, 12, 12, 12, 18], {k: INT for k in range(2, 9)}, total_cols=list(range(2, 9)))
    rr = []
    for kind, classes in (("intang", M.INTANG_CLASSES), ("tang", M.TANG_CLASSES)):
        mv = M.fixed_assets(kind)[2025]
        for k, c in enumerate(classes):
            rr.append([c[0], mv[k]["cost_open"], mv[k]["capex"], mv[k]["cost_close"], mv[k]["fund_close"], mv[k]["net_close"], c[5]])
    add_sheet(wb, "Riepilogo per classe", ["Classe", "Costo storico 1/1/2025", "Acquisizioni 2025", "Costo storico 31/12/2025", "Fondo 31/12/2025", "Valore netto 31/12/2025", "Aliquota"], rr, [60, 22, 18, 24, 20, 22, 10], {k: INT for k in range(1, 6)}, total_cols=[1, 2, 3, 4, 5],
              note="Riconciliazione con lo stato patrimoniale: i totali per classe coincidono con la nota integrativa del bilancio 2025.", freeze="A4")
    wb.save(path)
    return path


# ================================================================================================ bilancio di verifica e conto economico mensile
def verifica(path: Path) -> Path:
    wb = new_book()
    for y in (2025, 2024, 2023):
        pm = LG.pl_monthly(y)
        rows = []
        for k, a in enumerate(pm):
            rows.append([a["code"], a["name"], "Ricavo" if a["side"] == "R" else "Costo"] + [euro(v) for v in a["months"]] + [f"=SUM(D{k + 2}:O{k + 2})"])
        add_sheet(wb, f"CE mensile {y}", ["Codice", "Conto", "Natura"] + MS + ["Totale anno"], rows, [10, 66, 9] + [12] * 12 + [16], {k: MONEY for k in range(3, 16)}, total_cols=[], freeze="D2")
    for y in (2025, 2024, 2023):
        tb = LG.trial_balance(y)
        rows = [[r["code"], r["name"], euro(r["opening"]), euro(r["dare"]), euro(r["avere"]), euro(r["closing"])] for r in tb]
        add_sheet(wb, f"Bilancio di verifica {y}", ["Codice", "Conto", "Saldo iniziale", "Movimenti dare", "Movimenti avere", "Saldo finale"], rows, [10, 72, 18, 18, 18, 18], {k: MONEY for k in range(2, 6)}, total_cols=[2, 3, 4, 5],
                  total_label="Totale (il bilancio di verifica quadra)")
    # sintesi per voce
    rows = []
    for y in (2023, 2024, 2025):
        p = M.pl(y)
        rows.append([y, p["revenue"], p["a_total"], p["b_total"], p["ebit"], p["fin"], p["taxes"], p["net"]])
    add_sheet(wb, "Sintesi 2023-2025", ["Esercizio", "Ricavi", "Valore produzione", "Costi produzione", "EBIT", "Gestione finanziaria", "Imposte", "Utile"], rows, [12, 16, 18, 18, 14, 18, 14, 14], {k: INT for k in range(1, 8)})
    wb.save(path)
    return path


# ================================================================================================ fatture
def fatture_emesse(path: Path) -> Path:
    wb = new_book()
    for y in (2025, 2024, 2023):
        rows = []
        for k, inv in enumerate(LG.invoices_issued(y)):
            c = inv["customer"]
            rows.append([inv["number"], inv["date"], c["id"], c["name"], c["vat"], inv["desc"].capitalize(), inv["account"], euro(inv["taxable"]), c["iva"], euro(inv["vat"]), f"=H{k + 2}+J{k + 2}", inv["due"],
                         "Scissione dei pagamenti" if c["kind"] == "PA" else ("Non imponibile" if c["kind"] == "Estero" else "Ordinaria")])
        add_sheet(wb, f"Fatture emesse {y}", ["Numero", "Data", "Codice cliente", "Cliente", "Partita IVA", "Descrizione", "Conto di ricavo", "Imponibile", "Aliquota IVA", "IVA", "Totale fattura", "Scadenza", "Trattamento IVA"], rows,
                  [12, 12, 12, 46, 16, 26, 14, 14, 16, 12, 14, 12, 24], {1: DATE, 7: MONEY, 9: MONEY, 10: MONEY, 11: DATE}, total_cols=[7, 9, 10], total_label=f"Totale {y}")
    rows = []
    for c, v in LG.customer_revenue(2025):
        rows.append([c["id"], c["name"], c["sector"], c["region"], c["kind"], euro(v)] + [euro(sum(i["taxable"] for i in LG.invoices_issued(yy) if i["customer"]["id"] == c["id"])) for yy in (2024, 2023)])
    add_sheet(wb, "Fatturato per cliente", ["Codice", "Cliente", "Settore", "Regione", "Tipologia", "Fatturato 2025", "Fatturato 2024", "Fatturato 2023"], rows, [10, 46, 28, 16, 12, 16, 16, 16], {k: MONEY for k in (5, 6, 7)}, total_cols=[5, 6, 7])
    rows = []
    for o in LG.open_receivables(2025):
        rows.append([o["number"], o["customer"]["name"], o["date"], o["due"], euro(o["total"]), euro(o["open"]), (date(2025, 12, 31) - o["due"]).days])
    add_sheet(wb, "Scadenzario crediti 31-12-2025", ["Fattura", "Cliente", "Data", "Scadenza", "Importo", "Residuo", "Giorni di scaduto (negativo = non scaduta)"], rows, [12, 46, 12, 12, 14, 14, 30], {2: DATE, 3: DATE, 4: MONEY, 5: MONEY}, total_cols=[4, 5])
    wb.save(path)
    return path


def fatture_ricevute(path: Path) -> Path:
    wb = new_book()
    for y in (2025, 2024):
        rows = []
        for k, inv in enumerate(LG.invoices_received(y)):
            s = inv["supplier"]
            rows.append([inv["protocol"], inv["date"], s["id"], s["name"], s["vat"], inv["desc"], inv["account"], euro(inv["taxable"]), "22%", euro(inv["vat"]), f"=H{k + 2}+J{k + 2}", inv["due"]])
        add_sheet(wb, f"Fatture ricevute {y}", ["Protocollo", "Data", "Codice fornitore", "Fornitore", "Partita IVA", "Descrizione", "Conto di costo", "Imponibile", "Aliquota IVA", "IVA", "Totale", "Scadenza"], rows,
                  [12, 12, 12, 46, 16, 66, 12, 14, 10, 12, 14, 12], {1: DATE, 7: MONEY, 9: MONEY, 10: MONEY, 11: DATE}, total_cols=[7, 9, 10], total_label=f"Totale {y}")
    agg: Dict[str, int] = {}
    for inv in LG.invoices_received(2025):
        agg[inv["supplier"]["id"]] = agg.get(inv["supplier"]["id"], 0) + inv["taxable"]
    sp = {s["id"]: s for s in LG.suppliers()}
    rows = [[k, sp[k]["name"], sp[k]["category"], euro(v)] for k, v in sorted(agg.items(), key=lambda t: -t[1])]
    add_sheet(wb, "Acquisti per fornitore 2025", ["Codice", "Fornitore", "Categoria", "Imponibile 2025"], rows, [10, 50, 34, 18], {3: MONEY}, total_cols=[3])
    rows = [[o["protocol"], o["supplier"]["name"], o["date"], o["due"], euro(o["total"]), euro(o["open"])] for o in LG.open_payables(2025)]
    add_sheet(wb, "Scadenzario debiti 31-12-2025", ["Protocollo", "Fornitore", "Data", "Scadenza", "Importo", "Residuo"], rows, [12, 50, 12, 12, 14, 14], {2: DATE, 3: DATE, 4: MONEY, 5: MONEY}, total_cols=[4, 5])
    wb.save(path)
    return path


# ================================================================================================ commesse e budget
def commesse(path: Path) -> Path:
    wb = new_book()
    r = LG.rng("commesse")
    cust = LG.customers()
    kinds = ["Applicativo web", "App mobile", "Integrazione gestionale", "Cruscotto dati", "Portale servizi", "Manutenzione evolutiva", "Consulenza sicurezza", "Formazione"]
    rows = []
    total_rev = M.REVENUE[3]
    n = 54
    revs = LG.alloc(total_rev * 100, [r.uniform(0.3, 3.0) for _ in range(n)])
    for k in range(n):
        c = r.choice(cust[:40])
        rev = revs[k]
        hours = max(40, int(rev / 100 / r.uniform(52, 66)))
        cost_h = r.uniform(33, 41)
        ext = int(rev * r.uniform(0.02, 0.12))
        start = date(2025, r.randint(1, 11), r.randint(1, 27))
        rows.append([f"PRJ-2025-{k + 1:03d}", c["name"], r.choice(kinds), start, start + timedelta(days=r.randint(40, 240)), hours, round(hours * cost_h, 2), euro(ext), euro(rev), f"=I{k + 2}-G{k + 2}-H{k + 2}", f"=J{k + 2}/I{k + 2}"])
    add_sheet(wb, "Commesse 2025", ["Codice", "Cliente", "Tipo", "Inizio", "Fine", "Ore consuntivate", "Costo del personale (EUR)", "Costi esterni (EUR)", "Ricavi (EUR)", "Margine (EUR)", "Margine %"], rows,
              [16, 46, 24, 12, 12, 16, 22, 18, 16, 16, 12], {3: DATE, 4: DATE, 5: INT, 6: MONEY, 7: MONEY, 8: MONEY, 9: MONEY, 10: PCT}, total_cols=[5, 6, 7, 8, 9], total_label="Totale commesse")
    # budget 2026 vs consuntivo gen-ago
    r2 = LG.rng("budget2026")
    pm = LG.pl_monthly(2025)
    rows = []
    for k, a in enumerate(pm):
        growth = 1.17 if a["side"] == "R" else (1.10 if a["section"] in ("B.7", "B.6") else (1.14 if a["section"].startswith("B.9") else 1.04))
        bud = [euro(round(v * growth * r2.uniform(0.98, 1.02))) for v in a["months"]]
        act = [round(b * r2.uniform(0.94, 1.07), 2) if m < 8 else None for m, b in enumerate(bud)]
        rows.append([a["code"], a["name"]] + bud + [f"=SUM(C{k + 2}:N{k + 2})"] + [x if x is not None else "" for x in act[:8]] + [f"=SUM(P{k + 2}:W{k + 2})", f"=SUM(C{k + 2}:J{k + 2})", f"=X{k + 2}-Y{k + 2}"])
    hdr = ["Codice", "Conto"] + [f"Budget {m}" for m in MS] + ["Budget anno"] + [f"Consuntivo {m}" for m in MS[:8]] + ["Consuntivo gen-ago", "Budget gen-ago", "Scostamento"]
    add_sheet(wb, "Budget 2026 vs consuntivo", hdr, rows, [10, 60] + [12] * 12 + [16] + [12] * 8 + [16, 16, 14], {k: MONEY for k in range(2, 27)}, freeze="C2", note="Budget 2026 e consuntivo gennaio-agosto 2026 per conto (i consuntivi sono dati di gestione provvisori).", start_row=1)
    wb.save(path)
    return path


# ================================================================================================ modello del business plan
def modello_bp(path: Path) -> Path:
    wb = new_book()
    mh = [f"{MS[m % 12]} {2026 + m // 12}" for m in range(36)]
    ass = [["Ricavi 2026", bp.REV_TARGET[2026]], ["Ricavi 2027", bp.REV_TARGET[2027]], ["Ricavi 2028", bp.REV_TARGET[2028]], ["Rivalutazione retribuzioni ogni gennaio", bp.WAGE_RISE], ["Premi, welfare e altri costi sul costo di base", bp.PERSONNEL_EXTRA],
           ["Quota capitalizzata dello sviluppo", bp.CAPITALIZED_SHARE], ["Nuovo finanziamento (marzo 2026)", bp.NEW_LOAN], ["Tasso finanziamento esistente", 0.0505], ["Tasso nuovo finanziamento", 0.0535], ["DSO 2026", bp.DSO[2026]], ["DSO 2027", bp.DSO[2027]],
           ["DSO 2028", bp.DSO[2028]], ["DPO", bp.DPO], ["Aliquota IRES", 0.24], ["Aliquota IRAP", 0.039]]
    for n, v in bp.EXT_PCT.items():
        for k, y in enumerate(bp.YEARS):
            ass.append([f"{n} {y} (% ricavi)", v[k]])
    add_sheet(wb, "Ipotesi", ["Ipotesi", "Valore"], ass, [60, 18], {1: '#,##0.0000'}, freeze="A2")
    rev = bp.revenue()
    rows = [[n] + rev[n] for n in bp.REV_LINES]
    rows.append(["Totale ricavi"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}5)" for c in range(2, 38)])
    add_sheet(wb, "Ricavi", ["Linea"] + mh, rows, [44] + [11] * 36, {k: INT for k in range(1, 37)}, freeze="B2")
    pl = bp.pl()
    ext = bp.external_costs()
    pers = bp.personnel_lines()
    lines = [("Ricavi delle vendite e delle prestazioni", pl["ricavi"]), ("Incrementi di immobilizzazioni per lavori interni", pl["lavori_interni"]), ("Contributi in conto esercizio", pl["contributi"]), ("Altri ricavi e proventi", pl["altri_ricavi"])]
    rows = [[a, *b] for a, b in lines]
    rows.append(["Valore della produzione"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}5)" for c in range(2, 38)])
    r0 = len(rows) + 2
    for n, s in ext.items():
        rows.append([n, *s])
    rows.append(["Oneri diversi di gestione", *pl["oneri_diversi"]])
    rows.append(["Svalutazione dei crediti", *pl["svalutazione_crediti"]])
    for n, s in pers.items():
        rows.append([n, *s])
    last_cost = len(rows) + 1
    rows.append(["Totale costi operativi (senza ammortamenti)"] + [f"=SUM({get_column_letter(c)}{r0}:{get_column_letter(c)}{last_cost})" for c in range(2, 38)])
    rows.append(["EBITDA"] + [f"={get_column_letter(c)}6-{get_column_letter(c)}{last_cost + 1}" for c in range(2, 38)])
    rows.append(["Ammortamenti", *pl["ammortamenti"]])
    rows.append(["EBIT"] + [f"={get_column_letter(c)}{last_cost + 2}-{get_column_letter(c)}{last_cost + 3}" for c in range(2, 38)])
    rows.append(["Oneri finanziari netti"] + [pl["oneri_finanziari"][c] - pl["proventi_finanziari"][c] for c in range(36)])
    rows.append(["Risultato prima delle imposte"] + [f"={get_column_letter(c)}{last_cost + 4}-{get_column_letter(c)}{last_cost + 5}" for c in range(2, 38)])
    rows.append(["Imposte"] + pl["imposte"])
    rows.append(["Utile netto"] + [f"={get_column_letter(c)}{last_cost + 6}-{get_column_letter(c)}{last_cost + 7}" for c in range(2, 38)])
    add_sheet(wb, "CE mensile", ["Voce"] + mh, rows, [48] + [11] * 36, {k: INT for k in range(1, 37)}, freeze="B2")
    st = bp.staff()
    rows = [[pr["id"], pr["role"], pr["area"]] + [round(x) for x in st["per_person"][pr["id"]]] for pr in st["people"]]
    add_sheet(wb, "Personale", ["Matricola", "Ruolo", "Area"] + mh, rows, [10, 36, 26] + [11] * 36, {k: INT for k in range(3, 39)}, total_cols=list(range(3, 39)), freeze="D2")
    cx = bp.capex()
    add_sheet(wb, "Investimenti", ["Investimento"] + mh, [[k, *v] for k, v in cx.items()], [52] + [11] * 36, {k: INT for k in range(1, 37)}, total_cols=list(range(1, 37)), freeze="B2")
    b = bp.balance_sheet()
    names = [("Immobilizzazioni immateriali", "intang"), ("Immobilizzazioni materiali", "tang"), ("Immobilizzazioni finanziarie", "deposits"), ("Lavori in corso", "wip"), ("Crediti verso clienti", "trade"), ("Crediti tributari", "tax_cr"),
             ("Crediti verso altri", "other_cr"), ("Ratei e risconti attivi", "prepaid"), ("Disponibilità liquide", "cash"), ("TOTALE ATTIVO", "assets"), ("Patrimonio netto", "pn"), ("TFR", "tfr"), ("Debiti verso banche", "loan"),
             ("Debiti verso fornitori", "suppliers"), ("Debiti tributari", "tax_debt"), ("Debiti previdenziali", "social"), ("Altri debiti", "other_debt"), ("Ratei e risconti passivi", "accrued"), ("TOTALE PASSIVO", "liabilities")]
    add_sheet(wb, "Stato patrimoniale", ["Voce (fine mese)"] + mh, [[l, *b[k]] for l, k in names], [40] + [11] * 36, {k: INT for k in range(1, 37)}, freeze="B2")
    cf = bp.cash_flow()
    cl = [("Utile netto", "utile"), ("Ammortamenti", "ammortamenti"), ("Variazione TFR", "d_tfr"), ("Variazione lavori in corso", "d_wip"), ("Variazione crediti verso clienti", "d_trade"), ("Variazione altri crediti e risconti", "d_other_assets"),
          ("Variazione debiti verso fornitori", "d_suppliers"), ("Variazione debiti tributari, previdenziali, altri e ratei", "d_tax_social_other"), ("Flusso operativo", "operating"), ("Investimenti", "investing"), ("Finanziamenti", "loan"),
          ("Dividendi", "dividends"), ("Variazione delle disponibilità liquide", "change"), ("Disponibilità a inizio mese", "open"), ("Disponibilità a fine mese", "close")]
    add_sheet(wb, "Rendiconto finanziario", ["Voce"] + mh, [[l, *cf[k]] for l, k in cl], [52] + [11] * 36, {k: INT for k in range(1, 37)}, freeze="B2")
    wb.save(path)
    return path
