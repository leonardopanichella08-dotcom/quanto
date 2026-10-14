"""Piccola cassetta degli attrezzi per fogli Excel veri: intestazioni, larghezze, formati, riquadri bloccati, totali con formule."""
from __future__ import annotations

from datetime import date, datetime
from typing import Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
TOT_FILL = PatternFill("solid", fgColor="E5E7EB")
BODY_FONT = Font(name="Calibri", size=10)
THIN = Side(style="thin", color="D1D5DB")
MONEY = '#,##0.00;[Red]-#,##0.00'
INT = '#,##0;[Red]-#,##0'
DATE = "DD/MM/YYYY"
PCT = "0.0%"


def new_book() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def add_sheet(wb: Workbook, title: str, headers: Sequence[str], rows: Sequence[Sequence], widths: Optional[Sequence[float]] = None, formats: Optional[Dict[int, str]] = None,
              total_cols: Sequence[int] = (), total_label: str = "Totale", freeze: Optional[str] = "A2", note: Optional[str] = None, start_row: int = 1):
    """Aggiunge un foglio. ``formats``: indice di colonna (0-based) → formato numerico; ``total_cols``: colonne con riga finale di SUM."""
    ws = wb.create_sheet(title[:31])
    r0 = start_row
    if note:
        ws.cell(row=r0, column=1, value=note).font = Font(italic=True, color="6B7280", size=9)
        r0 += 2
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 32
    first = r0 + 1
    for r, row in enumerate(rows, first):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY_FONT
            if formats and (c - 1) in formats and v not in ("", None):
                cell.number_format = formats[c - 1]
            elif isinstance(v, (date, datetime)):
                cell.number_format = DATE
    last = first + len(rows) - 1
    if total_cols and rows:
        tr = last + 1
        ws.cell(row=tr, column=1, value=total_label).font = Font(bold=True, size=10)
        for c in total_cols:
            col = get_column_letter(c + 1)
            cell = ws.cell(row=tr, column=c + 1, value=f"=SUM({col}{first}:{col}{last})")
            cell.font = Font(bold=True, size=10)
            cell.number_format = (formats or {}).get(c, INT)
        for c in range(1, len(headers) + 1):
            ws.cell(row=tr, column=c).fill = TOT_FILL
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    if freeze:
        ws.freeze_panes = f"A{first}" if freeze == "A2" else freeze
    ws.sheet_view.zoomScale = 100
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{last}"
    return ws


def euro(c: int) -> float:
    """Centesimi interi → euro con due decimali."""
    return round(c / 100, 2)
