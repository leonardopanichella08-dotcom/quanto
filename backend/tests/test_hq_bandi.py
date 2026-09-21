import base64
import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.core import hq
from app.core.requirements_extractor import category_scope, extract_more_rules, extract_requirements, split_sentences
from main import app
from tests.conftest import manager_token, seed_pattern_bank

client = TestClient(app)

BANDO_TEXT = """Art. 3 - Spese ammissibili
Sono ammissibili esclusivamente le spese per beni strumentali nuovi di fabbrica, interconnessi ai sistemi aziendali.
Non sono ammissibili le spese di consulenza e le spese di personale.
Il costo orario del personale non può essere superiore a 40,00 euro/ora.
Le spese di comunicazione non possono superare il 3% del totale del progetto.
E' obbligatoria la perizia tecnica asseverata per tutti gli investimenti.
Il rispetto del principio DNSH è condizione di ammissibilità.
Le spese sono ammissibili solo se pagate con bonifico tracciabile; i pagamenti in contanti non sono ammessi.
Il beneficiario deve mantenere il vincolo di destinazione per 5 anni dall'ultimazione dell'investimento.
Il contributo è pari al 40% delle spese ammissibili. Sono ammissibili le spese sostenute dal 01/03/2026 fino al 31/12/2027.
Il soggetto proponente deve presentare la dichiarazione sostitutiva di atto notorio prima della liquidazione.
"""




def hq_headers():
    return {"Authorization": f"Bearer {manager_token()}"}


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def make_pdf(lines):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for ln in lines:
        c.drawString(40, y, ln[:110])
        y -= 16
    c.save()
    return buf.getvalue()


# ------------------------------------------------------------------ accesso HQ








def test_failed_logins_are_recorded_without_the_attempted_password():
    client.post("/api/v2/auth/login", json={"email": "ignoto@example.test", "password": "password-tentata-9"})
    ev = client.get("/api/v2/hq/timeline", params={"op": "auth.login"}, headers=hq_headers()).json()
    assert any(e["status"] == "DENIED" for e in ev)
    assert "password-tentata-9" not in str(ev)


# ------------------------------------------------------------------ memoria e timeline
def _demo(bando="QUANTO-SANDBOX-60", mode="realistic"):
    r = client.get("/api/v2/budget/demo", params={"bando_id": bando, "mode": mode})
    assert r.status_code == 200, r.text
    return r.json()


def test_validation_is_remembered_with_trace_and_replayable():
    scenario = _demo(mode="stress")
    v = client.post("/api/v2/budget/validate", json=scenario)
    assert v.status_code == 200
    body = v.json()
    assert body["run_id"] and len(body["trace"]["steps"]) > 300 and len(body["items"]) == 46
    h = hq_headers()
    run = client.get(f"/api/v2/hq/runs/{body['run_id']}", headers=h).json()
    assert run["response"]["merkle_root"] == body["merkle_root"] and run["request"]["project_id"] == scenario["project_id"]
    assert len(run["response"]["trace"]["steps"]) == len(body["trace"]["steps"])
    tl = client.get("/api/v2/hq/timeline", params={"op": "budget.validate"}, headers=h).json()
    assert tl and tl[0]["run_id"] == body["run_id"] and tl[0]["details"]["conformity_score"] == body["conformity_score"]
    assert client.get("/api/v2/hq/runs/999999", headers=h).status_code == 404


def test_identical_revalidation_does_not_flood_the_memory():
    scenario = _demo()
    a = client.post("/api/v2/budget/validate", json=scenario).json()
    b = client.post("/api/v2/budget/validate", json=scenario).json()
    assert a["run_id"] == b["run_id"]
    events = client.get("/api/v2/hq/timeline", params={"op": "budget.validate"}, headers=hq_headers()).json()
    assert len(events) == 1
    scenario["cost_items"][0]["fte_allocation"] = 0.4               # un cambiamento reale crea una nuova esecuzione
    c = client.post("/api/v2/budget/validate", json=scenario).json()
    assert c["run_id"] != a["run_id"]


def test_runs_store_no_personal_data_in_descriptions():
    scenario = _demo()
    scenario["cost_items"][0]["description"] = "Bonifico a RSSMRA80A01H501U IBAN IT60X0542811101000000123456"
    v = client.post("/api/v2/budget/validate", json=scenario).json()
    run = client.get(f"/api/v2/hq/runs/{v['run_id']}", headers=hq_headers()).json()
    stored = run["request"]["cost_items"][0]["description"]
    assert "RSSMRA80A01H501U" not in stored and "IT60X0542811101000000123456" not in stored and "TOK-" in stored


def test_project_dossier_collects_everything_about_a_project():
    scenario = _demo()
    pid = scenario["project_id"]
    v = client.post("/api/v2/budget/validate", json=scenario).json()
    client.post("/api/v2/registry/register", json={"project_id": pid, "merkle_root": v["merkle_root"]})
    client.get(f"/api/v2/registry/verify/{pid}", params={"merkle_root": v["merkle_root"]})
    client.post("/api/v2/budget/export/pdf", json=scenario)
    dossier = client.get(f"/api/v2/hq/projects/{pid}", headers=hq_headers()).json()
    ops = [e["op"] for e in dossier["timeline"]]
    assert ops == ["budget.validate", "registry.register", "registry.verify_root", "budget.export_pdf"]      # in ordine cronologico
    assert dossier["registration"]["merkle_root"] == v["merkle_root"] and len(dossier["runs"]) == 1
    assert dossier["documents"][0]["kind"] == "EXPORT_PDF" and len(dossier["documents"][0]["sha256"]) == 64
    assert any(p["project_id"] == pid for p in client.get("/api/v2/hq/projects", headers=hq_headers()).json())


def test_overview_operations_map_and_stats():
    client.post("/api/v2/budget/validate", json=_demo())
    ov = client.get("/api/v2/hq/overview", headers=hq_headers()).json()
    assert ov["counts"]["runs"] == 1 and ov["counts"]["bandi"] == 6 and ov["validations"] == 1 and ov["storage"]["engine"] == "PostgreSQL"
    assert ov["registry"]["intact"] is True and len(ov["recent_events"]) >= 1
    ops = {o["id"]: o for o in client.get("/api/v2/hq/operations", headers=hq_headers()).json()}
    assert len(ops) >= 20 and ops["budget.validate"]["stats"]["count"] == 1 and ops["budget.validate"]["stages"]
    assert all(o["stages"] and o["description"] for o in ops.values())
    assert ops["allocation.optimize"]["stats"]["count"] == 0


def test_every_user_operation_leaves_a_timeline_event():
    seed_pattern_bank()
    scenario = _demo()
    v = client.post("/api/v2/budget/validate", json=scenario).json()
    client.post("/api/v2/registry/register", json={"project_id": scenario["project_id"], "merkle_root": v["merkle_root"]})
    client.post("/api/v2/pattern/match", json={"bando_category": "X", "draft_budget": {"personnel_pct": 0.5, "assets_pct": 0.2, "consulting_pct": 0.2, "overhead_pct": 0.1}})
    client.post("/api/v2/allocation/optimize", json={"historical_expenses": [{"item_id": "E1", "category": "PERSONNEL", "amount_eur": 1000}],
                                                     "available_funding_lines": [{"fund_id": "F", "allowed_categories": ["PERSONNEL"]}]})
    client.post("/api/v2/bandi/QUANTO-SANDBOX-60/select")
    client.get("/api/v2/budget/template.xlsx")
    seen = {e["op"] for e in client.get("/api/v2/hq/timeline", params={"limit": 200}, headers=hq_headers()).json()}
    assert {"budget.demo", "budget.validate", "registry.register", "pattern.match", "allocation.optimize", "bandi.select", "budget.template"} <= seen


def test_timeline_filters_and_pagination():
    seed_pattern_bank()
    h = hq_headers()
    for i in range(3):
        client.post("/api/v2/pattern/match", json={"bando_category": "X", "draft_budget": {"personnel_pct": 0.5 + i / 10, "assets_pct": 0.2, "consulting_pct": 0.2, "overhead_pct": 0.1}})
    allev = client.get("/api/v2/hq/timeline", params={"op": "pattern.match"}, headers=h).json()
    assert len(allev) == 3 and allev[0]["id"] > allev[-1]["id"]
    page = client.get("/api/v2/hq/timeline", params={"op": "pattern.match", "before_id": allev[0]["id"]}, headers=h).json()
    assert len(page) == 2
    assert client.get("/api/v2/hq/timeline", params={"status": "DENIED"}, headers=h).json() == []
    assert client.get("/api/v2/hq/timeline", params={"limit": 0}, headers=h).status_code == 422


# ------------------------------------------------------------------ database (sola lettura)
def test_db_explorer_is_read_only_and_whitelisted():
    client.post("/api/v2/budget/validate", json=_demo())
    h = hq_headers()
    tables = {t["name"]: t for t in client.get("/api/v2/hq/db/tables", headers=h).json()}
    assert {"anchors", "bandi", "rules", "events", "runs", "documents", "requirements"} <= set(tables) and tables["runs"]["rows"] == 1
    runs = client.get("/api/v2/hq/db/table/runs", headers=h).json()
    assert runs["total"] == 1 and "caratteri]" in runs["rows"][0]["response_json"]            # payload pesanti riassunti, non riversati
    assert client.get("/api/v2/hq/db/table/pg_tables", headers=h).status_code == 404
    assert client.get("/api/v2/hq/db/table/runs;DROP TABLE runs", headers=h).status_code == 404
    assert client.post("/api/v2/hq/db/table/runs", headers=h).status_code == 405
    assert client.get("/api/v2/hq/db/table/rules", params={"limit": 3, "offset": 2}, headers=h).json()["limit"] == 3


# ------------------------------------------------------------------ biblioteca dei bandi
def test_library_lists_curated_bandi_with_coverage():
    lst = client.get("/api/v2/bandi").json()
    ids = {b["bando_id"] for b in lst}
    assert {"IPERAMMORTAMENTO-2026", "NUOVA-SABATINI", "HORIZON-EUROPE-MGA", "TRANSIZIONE-5.0-2024-2025", "FNC3-2024", "QUANTO-SANDBOX-60"} <= ids
    for b in lst:
        assert sum(b["coverage"].values()) == 60
    assert next(b for b in lst if b["bando_id"] == "QUANTO-SANDBOX-60")["coverage"]["NON_ATTIVO"] == 0


def test_every_rule_has_a_source_and_confidence_and_gaps_are_declared():
    for bid in ("IPERAMMORTAMENTO-2026", "NUOVA-SABATINI", "HORIZON-EUROPE-MGA", "TRANSIZIONE-5.0-2024-2025", "FNC3-2024"):
        d = client.get(f"/api/v2/bandi/{bid}").json()
        assert d["sources"] and d["legal_refs"] and d["not_specified"] and d["requirements"]
        assert all(r["source_ref"] and r["confidence"] for r in d["rules"]), bid
        assert len(d["coverage"]) == 60 and d["grant_rules"]["bando_id"] == bid


def test_horizon_indirect_cost_rate_is_25_percent_verified_on_the_official_pdf():
    """Regressione: un riassunto automatico riportava 15%; il testo ufficiale (MGA V1.2) dice 25%."""
    d = client.get("/api/v2/bandi/HORIZON-EUROPE-MGA").json()
    rules = {r["key"]: r for r in d["rules"]}
    assert rules["overhead_flat_rate_pct"]["value"] == 0.25 and rules["overhead_flat_base"]["value"] == "DIRECT_EXCL_SUBCONTRACTING"
    assert rules["overhead_flat_rate_pct"]["confidence"] == "PRIMARIA" and "15%" in rules["overhead_flat_rate_pct"]["source_ref"]


def test_secondary_sources_are_never_presented_as_primary():
    d = client.get("/api/v2/bandi/IPERAMMORTAMENTO-2026").json()
    rules = {r["key"]: r for r in d["rules"]}
    assert rules["requires_eu_origin"]["confidence"] == "SECONDARIA" and rules["appraisal_threshold_eur"]["confidence"] == "INTERPRETAZIONE"
    assert rules["eligible_categories"]["value"] == ["CAPITAL_ASSETS"]
    assert d["status"] == "APERTO" and next(b for b in client.get("/api/v2/bandi").json() if b["bando_id"] == "TRANSIZIONE-5.0-2024-2025")["status"] == "CHIUSO"


def test_partial_bando_is_shown_as_partial_with_items_to_review():
    b = next(x for x in client.get("/api/v2/bandi").json() if x["bando_id"] == "FNC3-2024")
    assert b["extraction_status"] == "PARTIAL" and b["requirements_to_review"] >= 1 and b["rules_count"] == 1


def test_selecting_a_bando_returns_rules_and_logs_the_choice():
    r = client.post("/api/v2/bandi/NUOVA-SABATINI/select")
    assert r.status_code == 200 and r.json()["grant_rules"]["vat_never_eligible"] is True and r.json()["grant_rules"]["eligible_categories"] == ["CAPITAL_ASSETS"]
    assert client.post("/api/v2/bandi/NOPE/select").status_code == 404
    assert client.get("/api/v2/bandi/NOPE").status_code == 404
    ev = client.get("/api/v2/hq/timeline", params={"op": "bandi.select"}, headers=hq_headers()).json()
    assert ev[0]["bando_id"] == "NUOVA-SABATINI" and ev[0]["details"]["rule_version_hash"]


def test_curated_rules_drive_a_real_validation_of_the_stress_demo():
    """Il demo completo, validato con le regole vere della Sabatini: solo i beni passano, il resto è respinto con motivo."""
    scenario = _demo("NUOVA-SABATINI", "stress")
    v = client.post("/api/v2/budget/validate", json=scenario).json()
    by_cat = {}
    for i in v["items"]:
        by_cat.setdefault(i["category"], set()).add(i["status"])
    assert by_cat["PERSONNEL"] == {"REJECTED"} and by_cat["CONSULTING"] == {"REJECTED"} and by_cat["TRAINING"] == {"REJECTED"}
    a15 = next(i for i in v["items"] if i["item_id"] == "A-15")
    assert a15["status"] == "REJECTED" and "esclusa dal bando" in a15["rejection_reason"]          # capannone = immobile
    a04 = next(i for i in v["items"] if i["item_id"] == "A-04")
    assert a04["status"] == "REJECTED" and 19 in a04["criteria_failed"]                                # usato
    a01 = next(i for i in v["items"] if i["item_id"] == "A-01")
    assert a01["computed_cost_eur"] == a01["original_cost_eur"] - 39600.0 and 54 in a01["criteria_failed"]   # IVA sempre esclusa


def test_references_include_de_minimis():
    refs = client.get("/api/v2/bandi/references").json()
    assert refs[0]["id"] == "DE-MINIMIS-2023-2831" and "300.000" in refs[0]["summary"] and refs[0]["confidence"] == "SECONDARIA"


# ------------------------------------------------------------------ upload di un bando
def test_upload_text_extracts_rules_scope_and_requirements():
    r = client.post("/api/v2/bandi/upload", json={"name": "Bando di prova beni", "filename": "bando.txt", "text": BANDO_TEXT})
    assert r.status_code == 200, r.text
    body = r.json()
    rules = {x["key"]: x for x in body["detail"]["rules"]}
    assert body["bando_id"] == "CUSTOM-BANDO-DI-PROVA-BENI" and body["characters_read"] > 500
    assert rules["max_hourly_rate_personnel"]["value"] == 40 and rules["max_communication_pct"]["value"] == 0.03
    assert rules["requires_iot"]["value"] is True and rules["requires_new_asset"]["value"] is True and rules["requires_dnsh"]["value"] is True
    assert rules["appraisal_threshold_eur"]["value"] == 0 and rules["min_durability_months"]["value"] == 60
    assert rules["contribution_rate_pct"]["value"] == 0.4 and rules["eligibility_start"]["value"] == "2026-03-01" and rules["eligibility_end"]["value"] == "2027-12-31"
    assert rules["eligible_categories"]["value"] == ["CAPITAL_ASSETS"]          # «ammissibili esclusivamente le spese per beni strumentali»
    assert all(x["confidence"] == "PARSING" for x in rules.values())
    assert body["requirements_total"] >= 8 and body["detail"]["coverage_summary"]["REGOLA_DEL_BANDO"] >= 10


def test_upload_flags_unclassified_obligations_instead_of_ignoring_them():
    body = client.post("/api/v2/bandi/upload", json={"name": "Bando con clausola strana", "text": BANDO_TEXT}).json()
    review = [r for r in body["detail"]["requirements"] if r["kind"] == "DA_REVISIONARE"]
    assert body["requirements_to_review"] == len(review) >= 1
    assert any("atto notorio" in r["text"] for r in review)


def test_upload_is_remembered_as_document_and_event_and_source():
    body = client.post("/api/v2/bandi/upload", json={"name": "Bando memoria", "text": BANDO_TEXT, "filename": "avviso.txt"}).json()
    h = hq_headers()
    docs = client.get("/api/v2/hq/documents", params={"kind": "BANDO_TEXT"}, headers=h).json()
    assert docs[0]["name"] == "avviso.txt" and docs[0]["bando_id"] == body["bando_id"] and len(docs[0]["sha256"]) == 64
    dossier = client.get(f"/api/v2/hq/bandi/{body['bando_id']}", headers=h).json()
    assert dossier["usage"]["uploaded_sources"][0]["sha256"] == body["sha256"]
    assert {e["op"] for e in dossier["timeline"]} >= {"bando.upload"}


def test_upload_pdf_is_parsed():
    pdf = make_pdf(BANDO_TEXT.splitlines())
    r = client.post("/api/v2/bandi/upload", json={"name": "Bando in PDF", "filename": "bando.pdf", "content_base64": b64(pdf)})
    assert r.status_code == 200, r.text
    rules = {x["key"]: x["value"] for x in r.json()["detail"]["rules"]}
    assert rules["requires_dnsh"] is True and rules["max_hourly_rate_personnel"] == 40
    assert client.get("/api/v2/hq/documents", params={"kind": "BANDO_PDF"}, headers=hq_headers()).json()[0]["name"] == "bando.pdf"


def test_upload_rejects_bad_inputs():
    assert client.post("/api/v2/bandi/upload", json={"name": "abc", "text": "troppo corto"}).status_code == 422
    assert client.post("/api/v2/bandi/upload", json={"name": "Bando", "filename": "x.pdf", "content_base64": b64(b"%PDF-1.4 non e un pdf vero" * 10)}).status_code == 422
    assert client.post("/api/v2/bandi/upload", json={"name": "Bando", "content_base64": "@@non-base64@@"}).status_code == 400
    assert client.post("/api/v2/bandi/upload", json={"name": "Bando"}).status_code == 422
    assert client.post("/api/v2/bandi/upload", json={"name": "Bando", "text": BANDO_TEXT, "bando_id": "id con spazi"}).status_code == 422


def test_uploaded_bando_rules_feed_a_validation():
    body = client.post("/api/v2/bandi/upload", json={"name": "Bando validabile", "text": BANDO_TEXT}).json()
    scenario = _demo(body["bando_id"], "stress")
    v = client.post("/api/v2/budget/validate", json=scenario).json()
    assert v["bando_id"] == body["bando_id"] and len(v["items"]) == 46
    assert {i["status"] for i in v["items"] if i["category"] == "PERSONNEL"} == {"REJECTED"}            # personale escluso dal testo


# ------------------------------------------------------------------ estrattore (unit)
def test_sentence_splitting_and_requirement_mapping():
    assert len(split_sentences(BANDO_TEXT)) >= 9
    reqs = extract_requirements(BANDO_TEXT)
    by_topic = {r["topic"].split(" / ")[0]: r for r in reqs}
    assert by_topic["DNSH"]["criteria"] == [28] or 28 in by_topic["DNSH"]["criteria"]
    assert any(r["kind"] == "DIVIETO" and 52 in r["criteria"] for r in reqs)                       # contanti vietati -> criterio 52
    assert any(r["kind"] == "LIMITE" and 37 in r["criteria"] for r in reqs)


def test_category_scope_inference():
    assert category_scope("Sono ammissibili solo le spese per macchinari e attrezzature.") == ["CAPITAL_ASSETS"]
    assert category_scope("Non sono ammissibili le spese di consulenza.") == ["CAPITAL_ASSETS", "OVERHEAD", "PERSONNEL", "TRAINING"]
    assert category_scope("Il progetto deve essere concluso entro 12 mesi.") is None


def test_flat_rate_base_is_inferred_from_the_wording():
    a = extract_more_rules("Le spese generali sono riconosciute con un tasso forfettario del 15% dei costi di personale.")
    b = extract_more_rules("È applicato un tasso forfettario del 25% dei costi diretti ammissibili.")
    assert a["overhead_flat_rate_pct"] == "0.15" and a["overhead_flat_base"] == "PERSONNEL"
    assert b["overhead_flat_rate_pct"] == "0.25" and b["overhead_flat_base"] == "DIRECT_EXCL_SUBCONTRACTING"


# ------------------------------------------------------------------ template e import Excel/CSV
def test_field_catalog_endpoint_lists_every_field_with_criteria():
    d = client.get("/api/v2/budget/fields").json()
    assert d["categories"] and len(d["fields"]) >= 60 and all("group" in f and "criteria" in f for f in d["fields"])


def test_template_has_all_columns_examples_and_guide():
    r = client.get("/api/v2/budget/template.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Voci", "Guida ai campi"]
    header = [c.value for c in wb["Voci"][1]]
    assert "item_id" in header and "dnsh_compliant" in header and "milestone_id" in header and len(header) >= 60
    assert wb["Voci"].max_row >= 6 and wb["Guida ai campi"].max_row >= 60


def test_template_roundtrips_through_import_and_validation():
    tpl = client.get("/api/v2/budget/template.xlsx").content
    imp = client.post("/api/v2/budget/import", json={"filename": "template.xlsx", "content_base64": b64(tpl)})
    assert imp.status_code == 200, imp.text
    body = imp.json()
    assert len(body["items"]) == 4 and body["errors"] == [] and body["ignored_columns"] == []
    assert {i["category"] for i in body["items"]} == {"PERSONNEL", "CAPITAL_ASSETS", "CONSULTING", "OVERHEAD"}
    scenario = _demo()
    scenario["cost_items"] = body["items"]
    assert client.post("/api/v2/budget/validate", json=scenario).status_code == 200


def test_import_reports_errors_per_row_and_keeps_valid_rows():
    wb = Workbook()
    ws = wb.active
    ws.title = "Voci"
    ws.append(["item_id", "description", "category", "source_c_ref", "amount_eur", "vat_recoverable", "colonna_ignota"])
    ws.append(["OK-1", "Consulenza valida", "CONSULTING", "DOC-1", "1.500,50", "sì", "x"])
    ws.append(["BAD-1", "Categoria sbagliata", "MAGIA", "DOC-2", 100, None, None])
    ws.append(["BAD-2", "Importo mancante", "CONSULTING", "DOC-3", None, None, None])
    ws.append([None, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    body = client.post("/api/v2/budget/import", json={"filename": "voci.xlsx", "content_base64": b64(buf.getvalue())}).json()
    assert [i["item_id"] for i in body["items"]] == ["OK-1"] and body["items"][0]["amount_eur"] == 1500.5 and body["items"][0]["vat_recoverable"] is True
    assert [e["row"] for e in body["errors"]] == [3, 4] and "category" in body["errors"][0]["message"] and "amount_eur" in body["errors"][1]["message"]
    assert body["ignored_columns"] == ["colonna_ignota"]


def test_import_csv_with_semicolons_and_bad_files():
    csv_data = "item_id;description;category;source_c_ref;ral_eur;fte_allocation\nP1;PM;PERSONNEL;DOC;38000;0,5\n".encode("utf-8")
    body = client.post("/api/v2/budget/import", json={"filename": "voci.csv", "content_base64": b64(csv_data)}).json()
    assert body["items"][0]["fte_allocation"] == 0.5 and body["errors"] == []
    assert client.post("/api/v2/budget/import", json={"filename": "voci.txt", "content_base64": b64(b"x")}).status_code == 400
    assert client.post("/api/v2/budget/import", json={"filename": "voci.csv", "content_base64": b64(b"a;b\n1;2\n")}).status_code == 400      # niente item_id
    assert client.post("/api/v2/budget/import", json={"filename": "voci.csv", "content_base64": "@@"}).status_code == 400
    ev = client.get("/api/v2/hq/timeline", params={"op": "budget.import"}, headers=hq_headers()).json()
    assert ev and ev[-1]["details"]["imported"] == 1


def test_demo_endpoint_modes_and_errors():
    assert len(_demo(mode="stress")["cost_items"]) == 46 and len(_demo(mode="realistic")["cost_items"]) == 15
    assert client.get("/api/v2/budget/demo", params={"mode": "boh"}).status_code == 422
    assert client.get("/api/v2/budget/demo", params={"bando_id": "NOPE"}).status_code == 404
    assert _demo("HORIZON-EUROPE-MGA")["grant_rules"]["overhead_flat_rate_pct"] == 0.25
