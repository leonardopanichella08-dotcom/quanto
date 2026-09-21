import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.registry import Registry, attestation_dict, verify_attestation
from main import app
from tests.conftest import tamper_sql

client = TestClient(app)

RULES = {"bando_id": "TRANSIZIONE-5.0-2026", "bando_name": "Piano Transizione 5.0", "max_hourly_rate_personnel": 35.0,
         "max_consulting_percentage": 0.20, "max_overhead_percentage": 0.07, "rule_version_hash": "a8f3b129c9e840134012480a2"}
ITEMS = [
    {"item_id": "LINE-001", "description": "PM Junior", "category": "PERSONNEL", "source_c_ref": "DOC-PAYROLL-2026-08",
     "ccnl_code": "TERZO_SETTORE", "employee_level": "3", "ral_eur": 38000, "fte_allocation": 0.5, "duration_months": 12},
    {"item_id": "LINE-002", "description": "Architect", "category": "PERSONNEL", "source_c_ref": "DOC-PAYROLL-2026-09",
     "ccnl_code": "METALMECCANICA", "employee_level": "5", "ral_eur": 58000, "fte_allocation": 0.8, "duration_months": 12},
    {"item_id": "LINE-003", "description": "Energy Audit", "category": "CONSULTING", "source_c_ref": "DOC-FACT-2026-104", "amount_eur": 30000},
]


def body(project="PRJ-TEST-1", items=None):
    return {"project_id": project, "grant_rules": RULES, "cost_items": items if items is not None else [dict(i) for i in ITEMS]}


def validate(**kw):
    r = client.post("/api/v2/budget/validate", json=body(**kw))
    assert r.status_code == 200, r.text
    return r.json()


def register(project, root):
    return client.post("/api/v2/registry/register", json={"project_id": project, "merkle_root": root})


# ------------------------------------------------------------------ Missione Uno
def test_health_and_criteria_catalog():
    assert client.get("/api/v2/health").json()["status"] == "OPERATIONAL"
    crit = client.get("/api/v2/budget/criteria").json()
    assert crit["total_criteria"] == 60 and crit["implemented_count"] == 60 and len(crit["criteria"]) == 60


def test_validate_seed_budget_totals_and_score_bounds():
    j = validate()
    assert j["status"] == "VALIDATED" and 0 <= j["conformity_score"] < 100
    assert j["total_requested_eur"] == 121395.82 and j["total_approved_eur"] == 88853.37
    assert round(j["total_requested_eur"] - j["total_approved_eur"], 2) == j["total_rejected_eur"]
    assert j["merkle_root"].startswith("0x") and len(j["merkle_root"]) == 66 and j["cep_id"] == "CEP-" + j["merkle_root"][2:18].upper()
    assert {c["criterion"] for c in j["budget_checks"]} == {48, 49, 55, 56, 60}
    assert all(i["criteria_not_evaluated"] is not None for i in j["items"])


def test_validate_is_idempotent_same_root():
    assert validate()["merkle_root"] == validate()["merkle_root"]


def test_approved_never_exceeds_requested_regression():
    j = validate(items=[ITEMS[0]])
    assert j["conformity_score"] == 100 and j["total_approved_eur"] <= j["total_requested_eur"]


def test_validate_with_rich_fields_via_api():
    item = {**ITEMS[2], "expense_subtype": "AUDIT", "expense_date": "2026-05-01", "payment_method": "BANK_TRANSFER"}
    j = validate(items=[ITEMS[0], item])
    assert j["items"][1]["criteria_checked"] and 52 in j["items"][1]["criteria_checked"]


def test_validate_input_errors():
    assert client.post("/api/v2/budget/validate", json=body(items=[])).status_code == 400
    assert client.post("/api/v2/budget/validate", json=body(items=[dict(ITEMS[0]), dict(ITEMS[0])])).status_code == 422
    bad = dict(ITEMS[2]); bad.pop("amount_eur")
    assert client.post("/api/v2/budget/validate", json=body(items=[bad])).status_code == 422
    bad_rules = body(); bad_rules["grant_rules"] = {**RULES, "max_consulting_percentage": 0.6, "max_overhead_percentage": 0.5}
    assert client.post("/api/v2/budget/validate", json=bad_rules).status_code == 422


def test_response_exposes_no_pii_fields():
    text = client.post("/api/v2/budget/validate", json=body()).text.lower()
    for forbidden in ("iban", "tax_code", "fiscal_code", "employee_name"):
        assert forbidden not in text


# ------------------------------------------------------------------ registro firmato
def test_register_verify_and_tamper_detection():
    j = validate(project="PRJ-REG")
    r = register("PRJ-REG", j["merkle_root"])
    assert r.status_code == 201
    att = r.json()
    assert att["seq"] == 1 and att["prev_hash"] == "0" * 64 and att["cep_id"] == j["cep_id"] and att["signature"]

    ok = client.get("/api/v2/registry/verify/PRJ-REG", params={"merkle_root": j["merkle_root"]}).json()
    assert ok["is_valid_and_unaltered"] and ok["registration_found"] and ok["signature_valid"] and ok["chain_intact"]

    wrong = client.get("/api/v2/registry/verify/PRJ-REG", params={"merkle_root": "0x" + "ab" * 32}).json()
    assert not wrong["is_valid_and_unaltered"] and wrong["registered_merkle_root"] == j["merkle_root"]

    same = client.post("/api/v2/registry/verify/recompute", json={k: body("PRJ-REG")[k] for k in ("project_id", "grant_rules", "cost_items")}).json()
    assert same["is_valid_and_unaltered"] and same["recomputed_from_data"]

    tampered = [dict(i) for i in ITEMS]; tampered[0]["ral_eur"] = 38001
    t = client.post("/api/v2/registry/verify/recompute", json={"project_id": "PRJ-REG", "grant_rules": RULES, "cost_items": tampered}).json()
    assert not t["is_valid_and_unaltered"] and t["provided_merkle_root"] != t["registered_merkle_root"]


def test_verify_unregistered_project_is_never_valid():
    r = client.get("/api/v2/registry/verify/NEVER", params={"merkle_root": "0x" + "cd" * 32}).json()
    assert r["registration_found"] is False and r["is_valid_and_unaltered"] is False and r["registered_merkle_root"] == ""


def test_register_twice_conflicts_and_reports_same_root():
    j = validate(project="PRJ-TWICE")
    assert register("PRJ-TWICE", j["merkle_root"]).status_code == 201
    dup = register("PRJ-TWICE", j["merkle_root"])
    assert dup.status_code == 409 and dup.json()["detail"]["same_root"] is True
    other = register("PRJ-TWICE", "0x" + "11" * 32)
    assert other.status_code == 409 and other.json()["detail"]["same_root"] is False


@pytest.mark.parametrize("root", ["0x123", "zz" * 32, ""])
def test_register_rejects_malformed_root(root):
    assert register("P", root).status_code == 422


def test_zero_root_is_refused():
    assert register("P", "0x" + "00" * 32).status_code == 400


def test_chain_links_entries_and_is_tamper_evident():
    for i in range(3):
        assert register(f"PRJ-CHAIN-{i}", "0x" + f"{i + 1:02x}" * 32).status_code == 201
    entries = Registry.entries()
    assert [e.seq for e in entries] == [1, 2, 3]
    assert entries[1].prev_hash == entries[0].entry_hash and entries[2].prev_hash == entries[1].entry_hash
    status = client.get("/api/v2/registry/status").json()
    assert status["intact"] and status["entries"] == 3 and status["head_hash"] == entries[2].entry_hash

    tamper_sql("UPDATE anchors SET merkle_root = %s WHERE seq = 2", ("ee" * 32,))   # manomissione diretta del DB
    broken = client.get("/api/v2/registry/status").json()
    assert broken["intact"] is False and broken["broken_at_seq"] == 2
    res = client.get("/api/v2/registry/verify/PRJ-CHAIN-1", params={"merkle_root": "0x" + "ee" * 32}).json()
    assert res["is_valid_and_unaltered"] is False and res["chain_intact"] is False


def test_deleting_an_entry_breaks_the_chain():
    for i in range(3):
        register(f"PRJ-DEL-{i}", "0x" + f"{i + 1:02x}" * 32)
    tamper_sql("DELETE FROM anchors WHERE seq = 2")
    assert client.get("/api/v2/registry/status").json()["intact"] is False


def test_attestation_verifies_offline_and_rejects_forgery():
    j = validate(project="PRJ-OFFLINE")
    att = register("PRJ-OFFLINE", j["merkle_root"]).json()
    assert client.get("/api/v2/registry/attestation/PRJ-OFFLINE").json()["signature"] == att["signature"]
    assert client.post("/api/v2/registry/verify/attestation", json=att).json()["valid"] is True
    forged = {**att, "merkle_root": "0x" + "99" * 32}
    assert client.post("/api/v2/registry/verify/attestation", json=forged).json()["valid"] is False
    assert client.get("/api/v2/registry/attestation/MISSING").status_code == 404


def test_signature_from_untrusted_key_is_rejected(monkeypatch):
    register("PRJ-KEY", "0x" + "12" * 32)
    monkeypatch.setenv("QUANTO_SIGNING_KEY", "ab" * 32)          # il server ora si fida di un'altra chiave
    assert Registry.verify_chain().intact is False
    att = Registry.lookup("PRJ-KEY")
    assert verify_attestation(attestation_dict(att)) is False
    monkeypatch.setenv("QUANTO_TRUSTED_PUBLIC_KEYS", att.public_key)   # rotazione: la chiave vecchia resta di fiducia
    assert Registry.verify_chain().intact is True


def test_invalid_signing_key_is_a_clear_error(monkeypatch):
    monkeypatch.setenv("QUANTO_SIGNING_KEY", "not-a-key")
    with pytest.raises(ValueError):
        Registry.register("P", "0x" + "12" * 32)


def test_public_key_endpoint_flags_dev_key(monkeypatch):
    assert client.get("/api/v2/registry/public-key").json()["is_dev_key"] is True
    monkeypatch.setenv("QUANTO_SIGNING_KEY", "cd" * 32)
    j = client.get("/api/v2/registry/public-key").json()
    assert j["is_dev_key"] is False and len(j["public_key"]) == 64


# ------------------------------------------------------------------ export
def test_xlsx_export_contains_cep_numeric_cells_and_qr():
    r = client.post("/api/v2/budget/export/xlsx", json=body(project="PRJ-XLSX"))
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["B3"].value.startswith("CEP-") and ws["B4"].value.startswith("0x")
    assert ws["F9"].value == 26282.70 and len(ws._images) == 1


def test_pdf_export_is_a_valid_pdf_with_cep_id():
    r = client.post("/api/v2/budget/export/pdf", json=body(project="PRJ-PDF"))
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF") and len(r.content) > 2000
    assert "CEP-" in r.headers["content-disposition"]
    assert client.post("/api/v2/budget/export/pdf", json=body(items=[])).status_code == 400


# ------------------------------------------------------------------ pattern / allocazione
def test_pattern_and_allocation_endpoints():
    p = client.post("/api/v2/pattern/match", json={"bando_category": "X", "draft_budget": {"personnel_pct": 0.58, "assets_pct": 0.12, "consulting_pct": 0.25, "overhead_pct": 0.05}})
    assert p.status_code == 200 and p.json()["main_deviation"]["category"] == "consulting_pct"
    assert client.post("/api/v2/pattern/match", json={"bando_category": "X", "draft_budget": {"foo": 1}}).status_code == 422
    assert client.post("/api/v2/pattern/match", json={"bando_category": "X", "draft_budget": {"personnel_pct": 0}}).status_code == 400
    a = client.post("/api/v2/allocation/optimize", json={
        "historical_expenses": [{"item_id": "E1", "category": "PERSONNEL", "amount_eur": 100000}],
        "available_funding_lines": [{"fund_id": "F", "allowed_categories": ["PERSONNEL"], "coverage_pct": 0.62}]})
    assert a.status_code == 200 and a.json()["net_cost_to_entity_eur"] == 38000.0 and a.json()["status"] == "OPTIMIZED"
    assert client.post("/api/v2/allocation/optimize", json={"historical_expenses": [], "available_funding_lines": []}).status_code == 400


def test_cors_is_not_wildcard():
    r = client.options("/api/v2/budget/validate", headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "POST"})
    assert "access-control-allow-origin" not in r.headers
